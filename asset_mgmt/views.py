from __future__ import annotations

import csv
import io
import json
from datetime import datetime, time, timedelta
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Count, F, Q, Sum
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import CreateView, DetailView, ListView, TemplateView, UpdateView

from .forms import (
    ApprovalRemarksForm, AssetBOMForm, AssetChangeRequestForm, AssetConditionAssessmentForm, AssetCriticalityAssessmentForm, AssetDisposalForm, AssetDocumentForm, AssetForm, AssetImageForm, AssetImportForm, AssetMeterForm, AssetMeterReadingForm, AssetStatusChangeForm, AssetTransferRequestForm, CalibrationPlanForm, CalibrationRecordForm,
    CAPAForm, CAPEXProposalForm, DowntimeForm, FishboneCauseForm, GoodsReceiptForm,
    InspectionPlanForm, InspectionRecordForm, LLFForm, MaintenanceRequestForm, MaterialIssueForm, MeterReadingForm,
    ModificationOrderForm, PermitToWorkForm, PMPlanForm, PMTaskForm, PurchaseOrderForm,
    PurchaseRequestForm, RCAForm, RejectionForm, ShutdownJobForm, ShutdownPlanForm,
    SparePartForm, StockTransferForm, UtilityMeterForm, VendorInvoiceForm, WorkOrderForm,
    WorkOrderOperationForm, WorkOrderSpareForm, WorkOrderLabourForm, RiskAssessmentForm,
    MaintenanceRequestItemForm, MaintenanceCatalogCodeForm, WorkOrderConfirmationForm, SettlementRuleForm,
    PMPlanCounterForm, MaintenanceTaskListForm, MaintenanceTaskListOperationForm, MaintenanceTaskListMaterialForm,
    MaintenanceStrategyForm, MaintenanceStrategyPackageForm, PMPlanStrategyPackageForm, ConditionRuleForm, WorkCentreCapacityForm, ServiceEntrySheetForm,
)
from .models import (
    Area, Asset, AssetBOM, AssetCategory, AssetSubcategory, AssetChangeRequest, AssetConditionAssessment, AssetCriticalityAssessment, AssetDisposalRequest, AssetDocument, AssetImage, AssetMeter, AssetMeterReading, AssetStatusHistory, AssetTransferRequest, AssetDowntimeEvent, AuditLog, CalibrationPlan, CalibrationRecord,
    CAPAAction, CAPEXProposal, Department, CostCentre, FishboneCause, FunctionalLocation, GoodsReceipt,
    InspectionPlan, InspectionRecord, InventoryBalance, LLFObservation, MaintenanceRequest, MaterialIssue,
    MeterReading, ModificationOrder, PermitToWork, Plant, PMCall, PMPlan, PMTask, PurchaseOrder,
    PurchaseRequest, PurchaseRequestApprovalHistory, RCARecord, ShutdownJob, ShutdownPlan,
    SparePart, StockTransaction, StockTransfer, StoreLocation, UtilityMeter, UserAssignment, VendorInvoice,
    WorkCentre, WorkOrder, WorkOrderOperation, WorkOrderSpare, WorkOrderLabour, RiskAssessment,
    MaintenanceRequestItem, MaintenanceCatalogCode, WorkOrderConfirmation, SettlementRule, SettlementPosting,
    InventoryReservation, PMPlanCounter, MaintenanceTaskList, MaintenanceTaskListOperation, MaintenanceTaskListMaterial,
    MaintenanceStrategy, MaintenanceStrategyPackage, PMPlanStrategyPackage, ConditionRule, ServiceEntrySheet,
)
from .permissions import assert_object_access, is_admin, is_hod_for, is_plant_head_for, scope_queryset_for_user, user_plant_ids
from .roles import ALL_ROLES, RoleRequiredMixin, role_can_access_url, user_role
from .services.audit import log_action, serialise_instance
from .services.availability import calculate_asset_availability
from .services.reliability import calculate_maintenance_cost_for_period, calculate_reliability_metrics
from .services.inventory import post_goods_receipt, post_material_issue, post_stock_transfer
from .services.pm import generate_pm_work_order
from .services.sap_pm import work_centre_capacity, pm_plan_is_due, confirm_operation
from .services.maintenance_workflow import transition_maintenance_request, transition_work_order, transition_permit, WO_TRANSITIONS
from .services.workflows import (
    approve_purchase_request, choose_hod, mark_pr_processing, reject_purchase_request,
    submit_purchase_request,
)


# -----------------------------------------------------------------------------
# Shared helpers
# -----------------------------------------------------------------------------

def accessible_plants(user):
    qs = Plant.objects.filter(active=True).select_related('company')
    ids = user_plant_ids(user)
    return qs if ids is None else qs.filter(id__in=ids)


def selected_plant(request):
    plants = accessible_plants(request.user)
    plant_id = request.GET.get('plant') or request.POST.get('plant')
    if plant_id:
        plant = plants.filter(pk=plant_id).first()
        if plant:
            return plant, plants
    return None, plants


def chart_json(payload):
    """Serialize dashboard data safely for inline JavaScript.

    Dashboard labels can include user-managed plant/category/location names. Escape
    HTML-significant characters so a label cannot terminate the surrounding script.
    """
    return (
        json.dumps(payload, ensure_ascii=False)
        .replace('&', '\\u0026')
        .replace('<', '\\u003c')
        .replace('>', '\\u003e')
        .replace('\u2028', '\\u2028')
        .replace('\u2029', '\\u2029')
    )


def role_dashboard_summary(role):
    summaries = {
        'Admin': ('Full system control', 'All plants, workflow, master data, approvals and audit visibility.'),
        'Asset Admin': ('Asset administration', 'Register, maintain, import, classify and audit the full Asset Master.'),
        'Production': ('Production issue reporting', 'Raise maintenance requests and monitor issue progress.'),
        'Maintenance': ('Engineering execution', 'Assigned work orders, PM, RCA, CAPA, LLF and MOC execution.'),
        'HOD': ('Engineering HOD review', 'Team workload, overdue jobs, approvals, shutdown and reliability KPIs.'),
        'Safety': ('Safety control', 'Permit, risk and inspection compliance.'),
        'Store': ('Engineering store', 'Location stock, receipts, issues, transfers and low-stock control.'),
        'Purchase': ('Materials management', 'Approved PR processing, PO creation and supplier progress.'),
        'Accounts': ('Invoice clearance', 'Invoice review, clearance and payment status.'),
        'Manager': ('Plant management', 'Plant-wide performance, availability, cost and approval review.'),
    }
    title, text = summaries.get(role, ('Role not assigned', 'Ask the administrator to assign a role and plant.'))
    return {'title': title, 'text': text}


def queryset_status_filter(qs, request):
    status = request.GET.get('status', '').strip()
    plant_id = request.GET.get('plant', '').strip()
    query = request.GET.get('q', '').strip()
    field_names = {field.name for field in qs.model._meta.get_fields()}
    if status and 'status' in field_names:
        qs = qs.filter(status=status)
    if request.GET.get('open') == '1' and 'status' in field_names:
        qs = qs.exclude(status__in=['CLOSED', 'CANCELLED'])
    if plant_id:
        if 'plant' in field_names:
            qs = qs.filter(plant_id=plant_id)
        elif 'asset' in field_names:
            qs = qs.filter(asset__plant_id=plant_id)
        elif 'work_order' in field_names:
            qs = qs.filter(work_order__plant_id=plant_id)
        elif 'meter' in field_names:
            qs = qs.filter(meter__plant_id=plant_id)
    if query:
        text_fields = [f.name for f in qs.model._meta.fields if f.get_internal_type() in ('CharField', 'TextField')]
        condition = Q()
        for field in text_fields[:12]:
            condition |= Q(**{f'{field}__icontains': query})
        if condition:
            qs = qs.filter(condition)
    return qs


def safe_pdf_response(title, lines, filename):
    """Generate a small approval-note PDF. Falls back to text if ReportLab is unavailable."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        pdf = canvas.Canvas(response, pagesize=A4)
        width, height = A4
        y = height - 50
        pdf.setTitle(title)
        pdf.setFont('Helvetica-Bold', 15)
        pdf.drawString(45, y, title)
        y -= 28
        pdf.setFont('Helvetica', 9)
        for label, value in lines:
            text = f'{label}: {value if value not in (None, "") else "-"}'
            chunks = [text[i:i + 105] for i in range(0, len(text), 105)] or ['']
            for chunk in chunks:
                if y < 45:
                    pdf.showPage()
                    pdf.setFont('Helvetica', 9)
                    y = height - 45
                pdf.drawString(45, y, chunk)
                y -= 14
            y -= 3
        pdf.showPage()
        pdf.save()
        return response
    except ImportError:
        body = title + '\n\n' + '\n'.join(f'{label}: {value}' for label, value in lines)
        return HttpResponse(body, content_type='text/plain')


class PlantScopedMixin:
    def get_queryset(self):
        return scope_queryset_for_user(super().get_queryset(), self.request.user)


class AuditCreateMixin:
    def form_valid(self, form):
        response = super().form_valid(form)
        log_action(self.request, 'CREATE', self.object)
        return response


class AuditUpdateMixin:
    previous_values = None

    def get_object(self, queryset=None):
        obj = super().get_object(queryset)
        assert_object_access(self.request.user, obj)
        self.previous_values = serialise_instance(obj)
        return obj

    def form_valid(self, form):
        response = super().form_valid(form)
        log_action(self.request, 'UPDATE', self.object, previous=self.previous_values)
        return response


class ModuleListView(RoleRequiredMixin, PlantScopedMixin, ListView):
    template_name = 'asset_mgmt/module_list.html'
    context_object_name = 'items'
    paginate_by = 50
    title = ''
    columns = []
    create_url_name = ''
    detail_url_name = ''
    update_url_name = ''

    def get_queryset(self):
        return queryset_status_filter(super().get_queryset(), self.request).select_related(*self.get_select_related())

    def get_select_related(self):
        return []

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({
            'title': self.title,
            'columns': self.columns,
            'create_url_name': self.create_url_name,
            'detail_url_name': self.detail_url_name,
            'update_url_name': self.update_url_name,
            'can_create': bool(self.create_url_name and role_can_access_url(self.request.user, self.create_url_name)),
            'can_update': bool(self.update_url_name and role_can_access_url(self.request.user, self.update_url_name)),
            'plants': accessible_plants(self.request.user),
            'selected_plant_id': self.request.GET.get('plant', ''),
        })
        return context


class SecureCreateView(AuditCreateMixin, RoleRequiredMixin, CreateView):
    template_name = 'asset_mgmt/form.html'
    title = ''
    success_url_name = 'dashboard'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def get_success_url(self):
        return reverse_lazy(self.success_url_name)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['title'] = self.title
        return context

    def form_invalid(self, form):
        messages.error(self.request, 'Please correct the highlighted fields.')
        return super().form_invalid(form)


class SecureUpdateView(AuditUpdateMixin, RoleRequiredMixin, UpdateView):
    template_name = 'asset_mgmt/form.html'
    title = ''
    success_url_name = 'dashboard'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def get_queryset(self):
        return scope_queryset_for_user(super().get_queryset(), self.request.user)

    def get_success_url(self):
        return reverse_lazy(self.success_url_name)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['title'] = self.title
        return context


# -----------------------------------------------------------------------------
# Dashboards (Observations 1, 2, 4, 16, 17, 19, 20, 23)
# -----------------------------------------------------------------------------

class DashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/dashboard.html'
    allowed_roles = ALL_ROLES

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated:
            role = user_role(request.user)
            if role == 'Asset Admin':
                return redirect('asset_dashboard')
            if role in {'HOD', 'Manager'}:
                return redirect('maintenance_manager_dashboard')
            if role == 'Maintenance':
                return redirect('engineering_dashboard')
            if role == 'Purchase':
                return redirect('mm_dashboard')
            if role == 'Store':
                return redirect('store_dashboard')
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        today = timezone.localdate()
        month_start = today.replace(day=1)

        assets = scope_queryset_for_user(Asset.objects.all(), self.request.user)
        requests = scope_queryset_for_user(MaintenanceRequest.objects.all(), self.request.user)
        work_orders = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user)
        pm_plans = scope_queryset_for_user(PMPlan.objects.all(), self.request.user)
        downtime = scope_queryset_for_user(AssetDowntimeEvent.objects.all(), self.request.user)
        spares = scope_queryset_for_user(SparePart.objects.all(), self.request.user)
        prs = scope_queryset_for_user(PurchaseRequest.objects.all(), self.request.user)
        pos = scope_queryset_for_user(PurchaseOrder.objects.all(), self.request.user)
        if plant:
            assets = assets.filter(plant=plant)
            requests = requests.filter(plant=plant)
            work_orders = work_orders.filter(plant=plant)
            pm_plans = pm_plans.filter(plant=plant)
            downtime = downtime.filter(asset__plant=plant)
            spares = spares.filter(plant=plant)
            prs = prs.filter(plant=plant)
            pos = pos.filter(plant=plant)

        period_start = timezone.make_aware(
            datetime.combine(month_start, time.min),
            timezone.get_current_timezone(),
        )
        period_end = timezone.now()
        reliability = calculate_reliability_metrics(
            assets=assets,
            work_orders=work_orders,
            downtime_events=downtime,
            period_start=period_start,
            period_end=period_end,
        )

        open_requests_qs = requests.exclude(status__in=['CLOSED', 'CANCELLED'])
        open_work_orders_qs = work_orders.exclude(status__in=['CLOSED', 'CANCELLED'])
        status_rows = list(assets.values('status').annotate(total=Count('id')).order_by('-total', 'status'))
        wo_rows = list(open_work_orders_qs.values('status').annotate(total=Count('id')).order_by('-total', 'status'))
        plant_rows = list(open_requests_qs.values('plant__code').annotate(total=Count('id')).order_by('-total', 'plant__code')[:12])
        downtime_period = downtime.filter(started_at__gte=period_start, started_at__lt=period_end)
        downtime_by_type = list(downtime_period.values('downtime_type').annotate(total=Count('id')).order_by('-total', 'downtime_type'))
        critical_q = Q(criticality__in=['HIGH', 'CRITICAL']) | Q(safety_critical=True)
        critical_availability = []
        for asset in assets.filter(critical_q).order_by('-criticality_score', 'asset_id')[:10]:
            metric = calculate_asset_availability(asset, period_start, period_end)
            critical_availability.append((asset.asset_id, metric['availability_percent']))

        context.update({
            'role': user_role(self.request.user), 'role_summary': role_dashboard_summary(user_role(self.request.user)),
            'plants': plants, 'selected_plant': plant,
            'total_assets': assets.count(),
            'critical_assets': assets.filter(critical_q).count(),
            'open_requests': open_requests_qs.count(),
            'open_work_orders': open_work_orders_qs.count(),
            'overdue_pm': pm_plans.filter(active=True, next_due_date__lt=today).count(),
            'low_stock_spares': spares.filter(current_stock__lte=F('minimum_stock')).count(),
            'monthly_mttr': reliability['mttr_hours'],
            'monthly_mtbf': reliability['mtbf_hours'],
            'monthly_mttr_display': f"{reliability['mttr_hours']:.2f}" if reliability['mttr_hours'] is not None else '-',
            'monthly_mtbf_display': f"{reliability['mtbf_hours']:.2f}" if reliability['mtbf_hours'] is not None else '-',
            'monthly_downtime': reliability['breakdown_downtime_hours'],
            'monthly_breakdowns': reliability['breakdown_failures'],
            'monthly_completed_breakdown_repairs': reliability['completed_breakdown_repairs'],
            'monthly_scheduled_hours': reliability['scheduled_hours'],
            'monthly_operating_hours': reliability['operating_hours'],
            'monthly_availability': reliability['availability_percent'],
            'monthly_failure_rate': reliability['failure_rate_per_1000_hours'],
            'reliability_warnings': reliability['warnings'],
            'open_purchase_requests': prs.exclude(status__in=['CLOSED', 'REJECTED', 'CANCELLED']).count(),
            'pending_pr_approvals': prs.filter(status='SUBMITTED', assigned_hod=self.request.user).count(),
            'open_purchase_orders': pos.exclude(status__in=['CLOSED', 'CANCELLED']).count(),
            'recent_work_orders': work_orders.select_related('asset', 'assigned_to', 'work_centre').order_by('-created_at')[:8],
            'recent_requests': requests.select_related('asset', 'work_centre').order_by('-created_at')[:8],
            'asset_status_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in status_rows], 'codes': [r['status'] for r in status_rows], 'values': [r['total'] for r in status_rows]}),
            'wo_status_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in wo_rows], 'codes': [r['status'] for r in wo_rows], 'values': [r['total'] for r in wo_rows]}),
            'plant_request_json': chart_json({'labels': [r['plant__code'] or 'Unassigned' for r in plant_rows], 'values': [r['total'] for r in plant_rows]}),
            'downtime_type_json': chart_json({'labels': [r['downtime_type'].replace('_', ' ').title() for r in downtime_by_type], 'values': [r['total'] for r in downtime_by_type]}),
            'critical_availability_json': chart_json({'labels': [row[0] for row in critical_availability], 'values': [row[1] for row in critical_availability]}),
            'dashboard_period_label': f'{month_start:%d %b %Y} – {today:%d %b %Y}',
        })
        return context


class AssetDashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/asset_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        assets = scope_queryset_for_user(Asset.objects.select_related('plant', 'area', 'functional_location', 'category'), self.request.user)
        if plant:
            assets = assets.filter(plant=plant)
        status_rows = list(assets.values('status').annotate(total=Count('id')).order_by('status'))
        lifecycle_rows = list(assets.values('lifecycle_status').annotate(total=Count('id')).order_by('lifecycle_status'))
        plant_rows = list(assets.values('plant__code').annotate(total=Count('id')).order_by('plant__code'))
        location_rows = list(assets.values('functional_location__code').annotate(total=Count('id')).order_by('-total')[:12])
        category_rows = list(assets.values('category__name').annotate(total=Count('id')).order_by('-total')[:12])
        critical_assets = assets.filter(Q(criticality__in=['HIGH', 'CRITICAL']) | Q(safety_critical=True))
        context.update({
            'plants': plants,
            'selected_plant': plant,
            'total_assets': assets.count(),
            'active_assets': assets.filter(is_active=True).count(),
            'critical_assets': critical_assets.count(),
            'running_assets': assets.filter(status='RUNNING').count(),
            'stopped_assets': assets.filter(status='STOPPED').count(),
            'breakdown_assets': assets.filter(status='BREAKDOWN').count(),
            'maintenance_assets': assets.filter(status='MAINTENANCE').count(),
            'asset_readiness_percent': round((assets.filter(status='RUNNING').count() / assets.filter(is_active=True).count()) * 100, 1) if assets.filter(is_active=True).exists() else 0.0,
            'recent_assets': assets.order_by('-created_at')[:10],
            'critical_asset_list': critical_assets.order_by('-criticality_score', 'asset_id')[:10],
            'asset_status_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in status_rows], 'codes': [r['status'] for r in status_rows], 'values': [r['total'] for r in status_rows]}),
            'lifecycle_json': chart_json({'labels': [r['lifecycle_status'].replace('_', ' ').title() for r in lifecycle_rows], 'values': [r['total'] for r in lifecycle_rows]}),
            'plant_asset_json': chart_json({'labels': [r['plant__code'] or 'Unassigned' for r in plant_rows], 'values': [r['total'] for r in plant_rows]}),
            'location_asset_json': chart_json({'labels': [r['functional_location__code'] or 'Unassigned' for r in location_rows], 'values': [r['total'] for r in location_rows]}),
            'category_asset_json': chart_json({'labels': [r['category__name'] or 'Unassigned' for r in category_rows], 'values': [r['total'] for r in category_rows]}),
        })
        return context


class MaintenanceManagerDashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/maintenance_manager_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        today = timezone.localdate()
        month_start = today.replace(day=1)
        tz = timezone.get_current_timezone()
        period_start = timezone.make_aware(datetime.combine(month_start, time.min), tz)
        period_end = timezone.now()
        assets = scope_queryset_for_user(Asset.objects.filter(is_active=True), self.request.user)
        work_orders = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user)
        downtime = scope_queryset_for_user(AssetDowntimeEvent.objects.all(), self.request.user)
        plans = scope_queryset_for_user(PMPlan.objects.filter(active=True), self.request.user).select_related('counter_meter').prefetch_related('counters__meter')
        work_centres = scope_queryset_for_user(WorkCentre.objects.filter(active=True), self.request.user)
        if plant:
            assets = assets.filter(plant=plant); work_orders = work_orders.filter(plant=plant)
            downtime = downtime.filter(asset__plant=plant); plans = plans.filter(plant=plant); work_centres = work_centres.filter(plant=plant)
        metrics = calculate_reliability_metrics(assets=assets, work_orders=work_orders, downtime_events=downtime, period_start=period_start, period_end=period_end)
        due_rows = []
        for plan in plans:
            due, reasons = pm_plan_is_due(plan, today)
            if due:
                due_rows.append((plan, '; '.join(reasons)))
        pm_calls = scope_queryset_for_user(PMCall.objects.filter(pm_plan__in=plans), self.request.user)
        persisted_calls = pm_calls.filter(call_date__gte=month_start, call_date__lte=today)
        persisted_call_count = persisted_calls.count()
        # A currently-due plan without a call in this reporting period is a missed/due call.
        # Do not let historical DUE/GENERATED rows from older periods hide it.
        due_without_call = sum(1 for plan, _ in due_rows if not persisted_calls.filter(pm_plan=plan).exists())
        pm_completed = persisted_calls.filter(status='COMPLETED').count()
        pm_due_count = persisted_call_count + due_without_call
        pm_remaining = max(pm_due_count - pm_completed, 0)
        compliance = round((pm_completed / pm_due_count) * 100, 1) if pm_due_count else None
        capacity_rows = []
        week_end = today + timedelta(days=6)
        for wc in work_centres:
            cap = work_centre_capacity(wc, today, week_end)
            capacity_rows.append({'work_centre': wc, **cap})
        capacity_rows.sort(key=lambda row: row['utilization_percent'], reverse=True)
        cost_breakdown = calculate_maintenance_cost_for_period(work_orders, period_start, period_end)
        maintenance_cost = cost_breakdown['total']
        pipeline_rows = list(
            work_orders.exclude(status__in=['CLOSED', 'CANCELLED'])
            .values('status').annotate(total=Count('id')).order_by('-total', 'status')
        )
        total_capacity = sum((Decimal(row['capacity_hours']) for row in capacity_rows), Decimal('0'))
        total_planned = sum((Decimal(row['planned_hours']) for row in capacity_rows), Decimal('0'))
        capacity_utilization = round(float((total_planned / total_capacity) * 100), 1) if total_capacity else 0.0
        context.update({
            'plants': plants, 'selected_plant': plant, 'metrics': metrics,
            'pm_compliance': compliance, 'pm_due': len(due_rows), 'pm_due_period': pm_due_count, 'pm_remaining': pm_remaining, 'pm_completed': pm_completed, 'due_pm_rows': due_rows[:12],
            'backlog_count': work_orders.exclude(status__in=['CLOSED','CANCELLED']).count(),
            'approval_queue': work_orders.filter(status='PENDING_APPROVAL').count(),
            'ready_schedule': work_orders.filter(status='READY_TO_SCHEDULE').count(),
            'scheduled_count': work_orders.filter(status__in=['SCHEDULED','DISPATCHED']).count(),
            'overdue_wo': work_orders.filter(planned_end__lt=timezone.now()).exclude(status__in=['COMPLETED','VERIFIED','TECO','COST_CLOSURE','CLOSED','CANCELLED']).count(),
            'maintenance_cost': maintenance_cost, 'cost_breakdown': cost_breakdown,
            'capacity_utilization': capacity_utilization, 'capacity_overloaded': sum(1 for row in capacity_rows if row['utilization_percent'] > 100),
            'open_breakdowns': work_orders.filter(work_type='BREAKDOWN').exclude(status__in=['CLOSED','CANCELLED']).count(),
            'capacity_rows': capacity_rows,
            'recent_breakdowns': work_orders.filter(work_type='BREAKDOWN').select_related('asset','work_centre','assigned_to').order_by('-created_at')[:10],
            'schedule_queue': work_orders.filter(status__in=['READY_TO_SCHEDULE','SCHEDULED','DISPATCHED']).select_related('asset','work_centre','assigned_to').order_by('planned_start')[:15],
            'manager_pipeline_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in pipeline_rows], 'codes': [r['status'] for r in pipeline_rows], 'values': [r['total'] for r in pipeline_rows]}),
            'manager_capacity_json': chart_json({'labels': [row['work_centre'].code for row in capacity_rows[:12]], 'values': [row['utilization_percent'] for row in capacity_rows[:12]]}),
            'manager_pm_json': chart_json({'labels': ['Completed', 'Due / Missed'], 'values': [pm_completed, pm_remaining]}),
            'manager_cost_json': chart_json({'labels': ['Materials', 'Labour', 'External Service', 'Other'], 'values': [float(cost_breakdown['material']), float(cost_breakdown['labour']), float(cost_breakdown['external_service']), float(cost_breakdown['other'])]}),
            'dashboard_period_label': f'{month_start:%d %b %Y} – {today:%d %b %Y}',
        })
        return context


class EngineeringDashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/engineering_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        work_orders = scope_queryset_for_user(WorkOrder.objects.select_related('asset', 'work_centre', 'assigned_to'), self.request.user)
        requests = scope_queryset_for_user(MaintenanceRequest.objects.select_related('asset', 'work_centre'), self.request.user)
        if plant:
            work_orders = work_orders.filter(plant=plant)
            requests = requests.filter(plant=plant)
        role = user_role(self.request.user)
        if role == 'Maintenance':
            work_orders = work_orders.filter(Q(assigned_to=self.request.user) | Q(created_by=self.request.user))
        pending_pr = scope_queryset_for_user(PurchaseRequest.objects.filter(status='SUBMITTED'), self.request.user)
        if role in ['HOD', 'Manager']:
            pending_pr = pending_pr.filter(Q(assigned_hod=self.request.user) | Q(plant__user_assignments__user=self.request.user, plant__user_assignments__designation__in=['HOD', 'PLANT_HEAD'])).distinct()
        else:
            pending_pr = pending_pr.none()
        status_rows = list(work_orders.values('status').annotate(total=Count('id')).order_by('status'))
        context.update({
            'plants': plants, 'selected_plant': plant,
            'assigned_work_orders': work_orders.exclude(status__in=['CLOSED', 'CANCELLED'])[:30],
            'overdue_work_orders': work_orders.filter(planned_end__lt=timezone.now()).exclude(status__in=['COMPLETED', 'VERIFIED', 'CLOSED', 'CANCELLED'])[:20],
            'open_requests': requests.exclude(status__in=['CLOSED', 'CANCELLED'])[:20],
            'pending_prs': pending_pr[:20],
            'wo_status_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in status_rows], 'values': [r['total'] for r in status_rows]}),
        })
        return context


class MMDashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/mm_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        prs = scope_queryset_for_user(PurchaseRequest.objects.all(), self.request.user)
        pos = scope_queryset_for_user(PurchaseOrder.objects.all(), self.request.user)
        receipts = scope_queryset_for_user(GoodsReceipt.objects.all(), self.request.user)
        if plant:
            prs, pos, receipts = prs.filter(plant=plant), pos.filter(plant=plant), receipts.filter(plant=plant)
        pr_rows = list(prs.values('status').annotate(total=Count('id')).order_by('status'))
        po_rows = list(pos.values('status').annotate(total=Count('id')).order_by('status'))
        context.update({
            'plants': plants, 'selected_plant': plant,
            'submitted_prs': prs.filter(status='SUBMITTED').count(), 'approved_prs': prs.filter(status='APPROVED').count(),
            'processing_prs': prs.filter(status='PROCESSING').count(), 'open_pos': pos.exclude(status__in=['CLOSED', 'CANCELLED']).count(),
            'pending_receipts': pos.filter(status__in=['RELEASED', 'PARTIAL_RECEIVED']).count(),
            'recent_prs': prs.select_related('item', 'work_centre')[:10], 'recent_pos': pos.select_related('item')[:10],
            'pr_status_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in pr_rows], 'values': [r['total'] for r in pr_rows]}),
            'po_status_json': chart_json({'labels': [r['status'].replace('_', ' ').title() for r in po_rows], 'values': [r['total'] for r in po_rows]}),
        })
        return context


class StoreDashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/store_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        items = scope_queryset_for_user(SparePart.objects.all(), self.request.user)
        balances = scope_queryset_for_user(InventoryBalance.objects.select_related('item', 'location'), self.request.user)
        transactions = scope_queryset_for_user(StockTransaction.objects.select_related('item', 'from_location', 'to_location'), self.request.user)
        if plant:
            items = items.filter(plant=plant)
            balances = balances.filter(item__plant=plant)
            transactions = transactions.filter(item__plant=plant)
        stock_value = sum((item.stock_value for item in items), Decimal('0'))
        context.update({
            'plants': plants, 'selected_plant': plant, 'stock_value': stock_value,
            'low_stock_count': items.filter(current_stock__lte=F('minimum_stock')).count(),
            'zero_stock_count': items.filter(current_stock=0).count(), 'total_items': items.count(),
            'low_stock_items': items.filter(current_stock__lte=F('minimum_stock'))[:15],
            'balances': balances[:20], 'transactions': transactions[:20],
        })
        return context


class UtilityDashboardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/utility_dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        readings = scope_queryset_for_user(MeterReading.objects.select_related('meter'), self.request.user)
        meters = scope_queryset_for_user(UtilityMeter.objects.all(), self.request.user)
        if plant:
            readings, meters = readings.filter(meter__plant=plant), meters.filter(plant=plant)
        recent = list(readings.order_by('-reading_date')[:60])
        totals = {}
        for reading in recent:
            label = reading.meter.get_utility_type_display()
            totals[label] = float(totals.get(label, 0)) + float(reading.consumption)
        context.update({
            'plants': plants, 'selected_plant': plant, 'meter_count': meters.count(),
            'recent_readings': recent[:20],
            'utility_json': chart_json({'labels': list(totals.keys()), 'values': list(totals.values())}),
        })
        return context


# -----------------------------------------------------------------------------
# Asset Master / Asset Register (PDF-focused update)
# -----------------------------------------------------------------------------

class AssetListView(ModuleListView):
    model = Asset
    template_name = 'asset_mgmt/asset_list.html'
    context_object_name = 'assets'
    title = 'Asset Master / Asset Register'
    columns = [('asset_id', 'Asset ID'), ('tag_number', 'Tag'), ('name', 'Name'), ('plant.code', 'Plant'), ('functional_location.code', 'Functional Location'), ('criticality', 'Criticality'), ('status', 'Status')]
    create_url_name = 'asset_create'; detail_url_name = 'asset_detail'; update_url_name = 'asset_update'
    paginate_by = 50

    def get_select_related(self):
        return ['company', 'plant', 'department', 'area', 'functional_location', 'category', 'subcategory', 'responsible_department', 'responsible_engineer']

    def get_queryset(self):
        qs = super().get_queryset()
        if self.request.GET.get('critical') == '1':
            qs = qs.filter(Q(criticality__in=['HIGH', 'CRITICAL']) | Q(safety_critical=True))
        for param, field in [('category', 'category_id'), ('department', 'department_id'), ('area', 'area_id'), ('lifecycle', 'lifecycle_status'), ('operational', 'status'), ('condition', 'current_condition')]:
            value = self.request.GET.get(param, '').strip()
            if value:
                qs = qs.filter(**{field: value})
        warranty = self.request.GET.get('warranty', '').strip()
        if warranty == 'active':
            today = timezone.localdate()
            qs = qs.filter(warranty_start_date__lte=today, warranty_end_date__gte=today)
        elif warranty == 'expired':
            qs = qs.filter(warranty_end_date__lt=timezone.localdate())
        if self.request.GET.get('safety_critical') == '1':
            qs = qs.filter(safety_critical=True)
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        base = scope_queryset_for_user(Asset.objects.all(), self.request.user)
        if self.request.GET.get('plant'):
            base = base.filter(plant_id=self.request.GET.get('plant'))
        pagination_params = self.request.GET.copy()
        pagination_params.pop('page', None)
        context.update({
            'total_assets': base.count(),
            'active_assets': base.filter(is_active=True).count(),
            'critical_assets': base.filter(Q(criticality__in=['HIGH', 'CRITICAL']) | Q(safety_critical=True)).count(),
            'warranty_expiring': base.filter(warranty_end_date__gte=timezone.localdate(), warranty_end_date__lte=timezone.localdate() + timedelta(days=60)).count(),
            'departments': Department.objects.filter(plant_id=self.request.GET.get('plant')) if self.request.GET.get('plant') else Department.objects.all()[:200],
            'areas': Area.objects.filter(plant_id=self.request.GET.get('plant')) if self.request.GET.get('plant') else Area.objects.all()[:200],
            'categories': AssetCategory.objects.filter(active=True).order_by('name'),
            'lifecycle_choices': Asset.LIFECYCLE_STATUS_CHOICES,
            'operational_choices': Asset.OPERATIONAL_STATUS_CHOICES,
            'condition_choices': Asset.CONDITION_CHOICES,
            'pagination_query': pagination_params.urlencode(),
        })
        return context


class AssetDetailView(RoleRequiredMixin, PlantScopedMixin, DetailView):
    model = Asset
    template_name = 'asset_mgmt/asset_detail.html'
    context_object_name = 'asset'

    def get_queryset(self):
        return super().get_queryset().select_related('company', 'plant', 'department', 'area', 'functional_location', 'category', 'subcategory', 'parent_asset', 'responsible_department', 'responsible_engineer', 'custodian', 'cost_centre')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        start = timezone.now() - timedelta(days=30)
        end = timezone.now()
        events = self.object.downtime_events.filter(started_at__lt=end, ended_at__gt=start)
        metric = calculate_asset_availability(self.object, start, end)
        context.update({
            'downtime_events': events[:20],
            'availability_30d': metric['availability_percent'],
            'downtime_30d': metric['downtime_hours'],
            'child_assets': self.object.sub_assets.select_related('plant', 'functional_location')[:100],
            'documents': self.object.documents.select_related('document_owner', 'approved_by')[:50],
            'images': self.object.images.select_related('uploaded_by')[:24],
            'bom_items': self.object.bom_items.select_related('spare_part')[:50],
            'meters': self.object.asset_meters.prefetch_related('readings')[:20],
            'status_history': self.object.status_history.select_related('changed_by')[:30],
            'condition_history': self.object.condition_assessments.select_related('assessor')[:20],
            'transfers': self.object.transfer_requests.select_related('destination_plant', 'destination_location')[:20],
            'change_requests': self.object.change_requests.select_related('requested_by', 'approved_by')[:20],
            'open_work_orders': self.object.work_orders.exclude(status__in=['CLOSED', 'CANCELLED'])[:20],
            'pm_plans': self.object.pm_plans.all()[:20],
            'disposal_requests': self.object.disposal_requests.all()[:10],
        })
        return context


class AssetCreateView(SecureCreateView):
    model = Asset
    form_class = AssetForm
    template_name = 'asset_mgmt/asset_form.html'
    title = 'Create Asset Master'
    success_url_name = 'asset_list'

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        form.instance.updated_by = self.request.user
        response = super().form_valid(form)
        AssetStatusHistory.objects.create(
            asset=self.object,
            new_lifecycle_status=self.object.lifecycle_status,
            new_operational_status=self.object.status,
            effective_date=self.object.status_effective_date or timezone.localdate(),
            reason='Asset master created',
            changed_by=self.request.user,
            reference='Asset creation',
        )
        return response


class AssetUpdateView(SecureUpdateView):
    model = Asset
    form_class = AssetForm
    template_name = 'asset_mgmt/asset_form.html'
    title = 'Update Asset Master'
    success_url_name = 'asset_list'

    def form_valid(self, form):
        form.instance.updated_by = self.request.user
        return super().form_valid(form)


class AssetDocumentUploadView(SecureCreateView):
    model = AssetDocument; form_class = AssetDocumentForm; title = 'Upload Asset Document'; success_url_name = 'asset_list'


class AssetImageUploadView(SecureCreateView):
    model = AssetImage; form_class = AssetImageForm; title = 'Upload Asset Image'; success_url_name = 'asset_list'
    def form_valid(self, form):
        form.instance.uploaded_by = self.request.user
        return super().form_valid(form)


class AssetMeterCreateView(SecureCreateView):
    model = AssetMeter; form_class = AssetMeterForm; title = 'Create Asset Meter'; success_url_name = 'asset_list'


class AssetMeterReadingCreateView(SecureCreateView):
    model = AssetMeterReading; form_class = AssetMeterReadingForm; title = 'Record Asset Meter Reading'; success_url_name = 'asset_list'
    def form_valid(self, form):
        form.instance.entered_by = self.request.user
        return super().form_valid(form)


class AssetTransferCreateView(SecureCreateView):
    model = AssetTransferRequest; form_class = AssetTransferRequestForm; title = 'Create Asset Transfer Request'; success_url_name = 'asset_list'
    def form_valid(self, form):
        form.instance.requested_by = self.request.user
        form.instance.current_plant = form.instance.asset.plant
        form.instance.current_location = form.instance.asset.functional_location
        form.instance.current_department = form.instance.asset.department
        return super().form_valid(form)


class AssetTransferDetailView(RoleRequiredMixin, PlantScopedMixin, DetailView):
    model = AssetTransferRequest; template_name = 'asset_mgmt/asset_transfer_detail.html'; context_object_name = 'transfer'


class AssetChangeRequestCreateView(SecureCreateView):
    model = AssetChangeRequest; form_class = AssetChangeRequestForm; title = 'Create Asset Change Request'; success_url_name = 'asset_list'
    def form_valid(self, form):
        form.instance.requested_by = self.request.user
        return super().form_valid(form)


@login_required
@require_POST
def asset_status_change(request, pk):
    if not role_can_access_url(request.user, 'asset_status_change'):
        raise PermissionDenied
    asset = get_object_or_404(scope_queryset_for_user(Asset.objects.all(), request.user), pk=pk)
    form = AssetStatusChangeForm(request.POST)
    if not form.is_valid():
        messages.error(request, '; '.join([err for field in form.errors.values() for err in field]))
        return redirect('asset_detail', pk=pk)
    previous_lifecycle = asset.lifecycle_status
    previous_status = asset.status
    new_lifecycle = form.cleaned_data['lifecycle_status'] or asset.lifecycle_status
    new_status = form.cleaned_data['operational_status'] or asset.status
    if asset.lifecycle_status in ['DISPOSED'] and not is_admin(request.user):
        messages.error(request, 'Disposed assets are locked. Admin correction is required.')
        return redirect('asset_detail', pk=pk)
    asset.lifecycle_status = new_lifecycle
    asset.status = new_status
    asset.status_effective_date = form.cleaned_data['effective_date']
    asset.status_reason = form.cleaned_data['reason']
    asset.status_changed_by = request.user
    asset.updated_by = request.user
    asset.full_clean()
    asset.save()
    AssetStatusHistory.objects.create(
        asset=asset,
        previous_lifecycle_status=previous_lifecycle,
        new_lifecycle_status=new_lifecycle,
        previous_operational_status=previous_status,
        new_operational_status=new_status,
        effective_date=form.cleaned_data['effective_date'],
        reason=form.cleaned_data['reason'],
        changed_by=request.user,
        reference='Manual status action',
    )
    log_action(request, 'STATUS_CHANGE', asset, previous={'lifecycle_status': previous_lifecycle, 'status': previous_status}, new={'lifecycle_status': new_lifecycle, 'status': new_status})
    messages.success(request, 'Asset status updated and history recorded.')
    return redirect('asset_detail', pk=pk)


@login_required
@require_POST
def asset_transfer_action(request, pk, action):
    if not role_can_access_url(request.user, 'asset_transfer_action'):
        raise PermissionDenied
    transfer = get_object_or_404(AssetTransferRequest.objects.select_related('asset'), pk=pk)
    assert_object_access(request.user, transfer.asset)
    if action == 'submit' and transfer.status == 'DRAFT':
        transfer.status = 'SUBMITTED'
        message = 'Transfer submitted for approval.'
    elif action == 'approve' and transfer.status == 'SUBMITTED':
        transfer.status = 'APPROVED'; transfer.approved_by = request.user
        message = 'Transfer approved.'
    elif action == 'dispatch' and transfer.status == 'APPROVED':
        transfer.status = 'DISPATCHED'
        message = 'Transfer dispatched.'
    elif action == 'receive' and transfer.status == 'DISPATCHED':
        transfer.complete_transfer(request.user)
        message = 'Transfer received and asset master location updated.'
        messages.success(request, message)
        return redirect('asset_detail', pk=transfer.asset.pk)
    elif action == 'reject' and transfer.status in ['SUBMITTED', 'APPROVED']:
        transfer.status = 'REJECTED'
        message = 'Transfer rejected.'
    else:
        messages.error(request, 'Invalid transfer action for the current status.')
        return redirect('asset_transfer_detail', pk=pk)
    transfer.save()
    log_action(request, 'TRANSFER_ACTION', transfer, remarks=f'Action: {action}')
    messages.success(request, message)
    return redirect('asset_transfer_detail', pk=pk)


@login_required
@require_POST
def asset_change_request_action(request, pk, action):
    if not role_can_access_url(request.user, 'asset_change_request_action'):
        raise PermissionDenied
    change = get_object_or_404(AssetChangeRequest.objects.select_related('asset'), pk=pk)
    assert_object_access(request.user, change.asset)
    if action == 'approve' and change.status in ['DRAFT', 'SUBMITTED']:
        change.status = 'APPROVED'; change.approved_by = request.user
        msg = 'Asset change request approved. Apply the approved value through controlled edit/action.'
    elif action == 'reject' and change.status in ['DRAFT', 'SUBMITTED']:
        change.status = 'REJECTED'; change.approved_by = request.user
        msg = 'Asset change request rejected.'
    else:
        messages.error(request, 'Invalid change-request action.')
        return redirect('asset_detail', pk=change.asset.pk)
    change.save()
    log_action(request, 'CHANGE_REQUEST_ACTION', change, remarks=f'Action: {action}')
    messages.success(request, msg)
    return redirect('asset_detail', pk=change.asset.pk)


@login_required
def asset_import_template(request):
    if not role_can_access_url(request.user, 'asset_import_template'):
        raise PermissionDenied
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="asset_master_import_template.csv"'
    writer = csv.writer(response)
    writer.writerow(['plant_code', 'department_code', 'area_code', 'functional_location_code', 'functional_location_name', 'asset_type', 'tag_number', 'name', 'description', 'category', 'subcategory', 'parent_asset_id', 'manufacturer', 'model_number', 'serial_number', 'purchase_date', 'commissioning_date', 'warranty_start_date', 'warranty_end_date', 'lifecycle_status', 'operational_status', 'criticality', 'safety_critical', 'responsible_department_code', 'cost_centre_code', 'purchase_cost'])
    writer.writerow(['PL01', 'MAINT', 'PROD', 'VAI-PL01-PROD-GF-N-PUMPHOUSE', 'Pump House', 'PUMP', 'P-101', 'Feed Pump P-101', 'Feed transfer pump', 'Rotating Equipment', 'Pump', '', 'ABC Pumps', 'PX100', 'SN001', '2024-01-10', '2024-02-10', '2024-02-10', '2025-02-09', 'ACTIVE', 'RUNNING', 'HIGH', 'false', 'MAINT', 'CC-MAINT', '250000'])
    return response


@login_required
def asset_import(request):
    if not role_can_access_url(request.user, 'asset_import'):
        raise PermissionDenied
    form = AssetImportForm(request.POST or None, request.FILES or None)
    result = None
    if request.method == 'POST' and form.is_valid():
        upload = form.cleaned_data['csv_file']
        if upload.name.lower().endswith('.xlsx'):
            messages.error(request, 'Use the command-line real-data importer for SAP-style Excel files: py manage.py import_real_engineering_data --dry-run')
            return redirect('asset_import')
        decoded = upload.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(decoded))
        errors, prepared = [], []
        allowed = accessible_plants(request.user)
        for line_no, row in enumerate(reader, start=2):
            try:
                plant = allowed.get(code=row['plant_code'].strip())
                area, _ = Area.objects.get_or_create(plant=plant, code=row['area_code'].strip() or 'GEN', defaults={'name': row.get('area_code', 'General')[:150]})
                dept = None
                if row.get('department_code'):
                    dept, _ = Department.objects.get_or_create(plant=plant, code=row['department_code'].strip(), defaults={'name': row['department_code'].strip()})
                cost_centre = None
                if row.get('cost_centre_code'):
                    cost_centre, _ = CostCentre.objects.get_or_create(plant=plant, code=row['cost_centre_code'].strip(), defaults={'name': row['cost_centre_code'].strip()})
                fl_code = row.get('functional_location_code', '').strip()
                if fl_code:
                    location = FunctionalLocation.objects.filter(code=fl_code, plant=plant).first()
                    if not location:
                        location = FunctionalLocation(plant=plant, area=area, code=fl_code, name=row.get('functional_location_name', fl_code)[:180])
                else:
                    location = FunctionalLocation(plant=plant, area=area, name=row['functional_location_name'].strip())
                category = None
                if row.get('category'):
                    category, _ = AssetCategory.objects.get_or_create(name=row['category'].strip()[:120])
                subcategory = None
                if category and row.get('subcategory'):
                    subcategory, _ = AssetSubcategory.objects.get_or_create(category=category, code=row['subcategory'].strip().upper()[:30], defaults={'name': row['subcategory'].strip()[:120]})
                asset = Asset(
                    plant=plant, department=dept, area=area, functional_location=location, cost_centre=cost_centre,
                    asset_type=row['asset_type'].strip().upper() or 'OTHER', tag_number=row['tag_number'].strip(), name=row['name'].strip(),
                    asset_description=row.get('description', '').strip(), category=category, subcategory=subcategory,
                    manufacturer=row.get('manufacturer', '').strip(), model_number=row.get('model_number', '').strip(), serial_number=row.get('serial_number', '').strip(),
                    purchase_date=datetime.strptime(row['purchase_date'], '%Y-%m-%d').date() if row.get('purchase_date') else None,
                    commissioning_date=datetime.strptime(row['commissioning_date'], '%Y-%m-%d').date() if row.get('commissioning_date') else None,
                    warranty_start_date=datetime.strptime(row['warranty_start_date'], '%Y-%m-%d').date() if row.get('warranty_start_date') else None,
                    warranty_end_date=datetime.strptime(row['warranty_end_date'], '%Y-%m-%d').date() if row.get('warranty_end_date') else None,
                    lifecycle_status=row.get('lifecycle_status', 'ACTIVE').strip().upper() or 'ACTIVE', status=row.get('operational_status', 'RUNNING').strip().upper() or 'RUNNING',
                    criticality=row.get('criticality', 'MEDIUM').strip().upper() or 'MEDIUM', safety_critical=row.get('safety_critical', '').strip().lower() in ['1', 'true', 'yes'],
                    purchase_cost=Decimal(row['purchase_cost']) if row.get('purchase_cost') else None,
                    created_by=request.user, updated_by=request.user,
                )
                if Asset.objects.filter(plant=plant, tag_number=asset.tag_number).exists():
                    raise ValidationError(f'Duplicate tag number {asset.tag_number} in plant {plant.code}.')
                location.full_clean()
                asset.full_clean(exclude=['asset_id', 'qr_token', 'qr_code_value'])
                prepared.append((location, asset))
            except Exception as exc:
                errors.append(f'Line {line_no}: {exc}')
        if not errors and not form.cleaned_data['validate_only']:
            with transaction.atomic():
                for location, asset in prepared:
                    if not location.pk:
                        location.save()
                    asset.functional_location = location
                    asset.save()
                    AssetStatusHistory.objects.create(asset=asset, new_lifecycle_status=asset.lifecycle_status, new_operational_status=asset.status, reason='Bulk asset import', changed_by=request.user)
                    log_action(request, 'CREATE', asset, remarks='Bulk CSV asset import')
            messages.success(request, f'{len(prepared)} assets imported successfully.')
            return redirect('asset_list')
        result = {'valid_count': len(prepared), 'errors': errors, 'validate_only': form.cleaned_data['validate_only']}
    return render(request, 'asset_mgmt/asset_import.html', {'form': form, 'result': result})


# -----------------------------------------------------------------------------
# Maintenance, PM, shutdown, RCA/CAPA/LLF/MOC
# -----------------------------------------------------------------------------


class AssetBOMCreateView(SecureCreateView):
    model = AssetBOM; form_class = AssetBOMForm; title = 'Add Asset BOM Component'; success_url_name = 'asset_list'
    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get('asset'): initial['asset'] = self.request.GET.get('asset')
        return initial
    def get_success_url(self): return reverse_lazy('asset_detail', kwargs={'pk': self.object.asset_id})


class AssetBOMUpdateView(SecureUpdateView):
    model = AssetBOM; form_class = AssetBOMForm; title = 'Update Asset BOM Component'; success_url_name = 'asset_list'
    def get_success_url(self): return reverse_lazy('asset_detail', kwargs={'pk': self.object.asset_id})


class AssetConditionAssessmentCreateView(SecureCreateView):
    model = AssetConditionAssessment; form_class = AssetConditionAssessmentForm; title = 'Asset Condition Assessment'; success_url_name = 'asset_list'
    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get('asset'): initial['asset'] = self.request.GET.get('asset')
        return initial
    def form_valid(self, form):
        form.instance.assessor = self.request.user
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('asset_detail', kwargs={'pk': self.object.asset_id})


class AssetCriticalityAssessmentCreateView(SecureCreateView):
    model = AssetCriticalityAssessment; form_class = AssetCriticalityAssessmentForm; title = 'Asset Criticality Assessment'; success_url_name = 'asset_list'
    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get('asset'): initial['asset'] = self.request.GET.get('asset')
        return initial
    def form_valid(self, form):
        form.instance.assessed_by = self.request.user
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('asset_detail', kwargs={'pk': self.object.asset_id})


class MaintenanceRequestListView(ModuleListView):
    model = MaintenanceRequest; title = 'Maintenance Requests'
    template_name = 'asset_mgmt/maintenance_request_list.html'; context_object_name = 'requests'
    columns = []
    create_url_name = 'maintenance_request_create'; update_url_name = 'maintenance_request_update'
    def get_select_related(self): return ['asset', 'plant', 'work_centre', 'reported_by', 'screened_by']


class MaintenanceRequestCreateView(SecureCreateView):
    model = MaintenanceRequest; form_class = MaintenanceRequestForm; title = 'Report Maintenance Requirement'; success_url_name = 'maintenance_request_list'
    template_name = 'asset_mgmt/maintenance_request_form.html'
    def form_valid(self, form):
        form.instance.reported_by = self.request.user
        form.instance.plant = form.instance.asset.plant
        form.instance.status = 'REPORTED'
        return super().form_valid(form)


class MaintenanceRequestUpdateView(SecureUpdateView):
    model = MaintenanceRequest; form_class = MaintenanceRequestForm; title = 'Update Maintenance Request'; success_url_name = 'maintenance_request_list'
    template_name = 'asset_mgmt/maintenance_request_form.html'
    def form_valid(self, form):
        form.instance.reported_by = self.object.reported_by
        form.instance.plant = form.instance.asset.plant
        form.instance.status = self.object.status
        return super().form_valid(form)


@login_required
@require_GET
def maintenance_request_options(request):
    """Return plant-scoped asset and work-centre choices for the request form."""
    if not role_can_access_url(request.user, 'maintenance_request_create'):
        raise PermissionDenied
    plant_id = request.GET.get('plant', '').strip()
    if not plant_id.isdigit():
        return JsonResponse({'assets': [], 'work_centres': []}, status=400)
    plant = get_object_or_404(accessible_plants(request.user), pk=plant_id)
    assets = Asset.objects.filter(plant=plant).order_by('asset_id')
    work_centres = WorkCentre.objects.filter(plant=plant).order_by('code')
    return JsonResponse({
        'assets': [{'id': asset.pk, 'label': str(asset)} for asset in assets],
        'work_centres': [{'id': centre.pk, 'label': str(centre)} for centre in work_centres],
    })


@login_required
@require_POST
def maintenance_request_action(request, pk, action):
    if not role_can_access_url(request.user, 'maintenance_request_action'):
        raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(MaintenanceRequest.objects.all(), request.user), pk=pk)
    role = user_role(request.user)
    screening_actions = {'start_screening', 'request_info', 'accept', 'reject'}
    if action in screening_actions and role not in {'Admin', 'Maintenance', 'HOD', 'Manager'}:
        raise PermissionDenied
    if action == 'resubmit' and role not in {'Admin', 'Production', 'Maintenance', 'HOD'}:
        raise PermissionDenied
    if action == 'close' and role not in {'Admin', 'Maintenance', 'HOD', 'Manager'}:
        raise PermissionDenied
    try:
        previous = serialise_instance(obj)
        transition_maintenance_request(obj, action, request.user, request.POST.get('notes', ''))
        log_action(request, f'MR_{action.upper()}', obj, previous=previous)
        messages.success(request, f'{obj.request_no} moved to {obj.get_status_display()}.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('maintenance_request_list')


class MaintenanceRequestItemCreateView(SecureCreateView):
    model = MaintenanceRequestItem; form_class = MaintenanceRequestItemForm; title = 'Add Notification Item'; success_url_name = 'maintenance_request_list'
    def dispatch(self, request, *args, **kwargs):
        self.notification = get_object_or_404(scope_queryset_for_user(MaintenanceRequest.objects.all(), request.user), pk=kwargs['pk'])
        return super().dispatch(request, *args, **kwargs)
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.request.method in {'POST', 'PUT'}:
            kwargs['instance'] = MaintenanceRequestItem(maintenance_request=self.notification)
        return kwargs
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        field_types = {'object_part':'OBJECT_PART', 'damage_code':'DAMAGE', 'cause_code':'CAUSE', 'activity_code':'ACTIVITY'}
        for field_name, catalog_type in field_types.items():
            if field_name in form.fields:
                form.fields[field_name].queryset = MaintenanceCatalogCode.objects.filter(
                    plant=self.notification.plant, catalog_type=catalog_type, active=True
                ).order_by('code_group','code')
        return form
    def form_valid(self, form):
        form.instance.maintenance_request = self.notification
        return super().form_valid(form)


class MaintenanceCatalogListView(ModuleListView):
    model = MaintenanceCatalogCode; title = 'Maintenance Failure Catalogs'
    columns = [('plant.code','Plant'),('catalog_type','Catalog'),('code_group','Group'),('code','Code'),('description','Description'),('active','Active')]
    create_url_name = 'maintenance_catalog_create'
    def get_select_related(self): return ['plant']


class MaintenanceCatalogCreateView(SecureCreateView):
    model = MaintenanceCatalogCode; form_class = MaintenanceCatalogCodeForm; title = 'Create Maintenance Catalog Code'; success_url_name = 'maintenance_catalog_list'


class WorkOrderListView(ModuleListView):
    model = WorkOrder; title = 'Work Orders'
    columns = [('wo_number', 'WO'), ('asset.asset_id', 'Asset'), ('plant.code', 'Plant'), ('work_centre.code', 'Work Centre'), ('work_type', 'Type'), ('activity_type', 'Activity'), ('shift', 'Shift'), ('assigned_to.username', 'Assigned To'), ('status', 'Status')]
    create_url_name = 'work_order_create'; detail_url_name = 'work_order_detail'; update_url_name = 'work_order_update'
    def get_select_related(self): return ['asset', 'plant', 'work_centre', 'assigned_to', 'supervisor']
    def get_queryset(self):
        qs = super().get_queryset()
        for param, field in [('work_type', 'work_type'), ('activity', 'activity_type'), ('shift', 'shift')]:
            value = self.request.GET.get(param, '').strip()
            if value:
                qs = qs.filter(**{field: value})
        return qs


class WorkOrderDetailView(RoleRequiredMixin, PlantScopedMixin, DetailView):
    model = WorkOrder; template_name = 'asset_mgmt/work_order_detail.html'; context_object_name = 'work_order'
    def get_queryset(self):
        return super().get_queryset().select_related('asset', 'plant', 'work_centre', 'assigned_to', 'supervisor', 'maintenance_request', 'verified_by')
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        wo = self.object
        role = user_role(self.request.user)
        candidate_actions = [action for action, (source, target) in WO_TRANSITIONS.items() if source == wo.status]
        def permitted(action):
            if action in {'approve', 'verify', 'teco', 'close'}:
                return role in {'Admin', 'HOD', 'Manager'}
            if action == 'cost_close':
                return role in {'Admin', 'HOD', 'Manager', 'Accounts'}
            return role in {'Admin', 'Maintenance', 'HOD', 'Manager'}
        context.update({
            'operations': wo.operations.select_related('work_centre', 'assigned_to'),
            'planned_spares': wo.spares_used.select_related('spare_part'),
            'labour_entries': wo.labour_entries.select_related('person'),
            'material_issues': wo.material_issues.select_related('item', 'source_location'),
            'permits': wo.permits.all(),
            'risk_assessments': wo.risk_assessments.all(),
            'downtime_events': wo.downtime_events.all(),
            'confirmations': wo.confirmations.select_related('operation', 'person'),
            'reservations': InventoryReservation.objects.filter(work_order_spare__work_order=wo).select_related('item', 'location', 'work_order_spare__operation'),
            'purchase_requests': wo.purchase_requests.select_related('item', 'source_operation'),
            'settlement_rules': wo.settlement_rules.all(),
            'settlement_postings': wo.settlement_postings.select_related('rule'),
            'material_cost': wo.material_cost,
            'labour_cost': wo.labour_cost,
            'calculated_actual_cost': wo.calculated_actual_cost,
            'workflow_actions': [action for action in candidate_actions if permitted(action)],
        })
        return context


class WorkOrderCreateView(SecureCreateView):
    model = WorkOrder; form_class = WorkOrderForm; title = 'Create Work Order'; success_url_name = 'work_order_list'
    def get_initial(self):
        initial = super().get_initial()
        mr_id = self.request.GET.get('maintenance_request')
        if mr_id:
            mr = scope_queryset_for_user(MaintenanceRequest.objects.all(), self.request.user).filter(pk=mr_id, status='ACCEPTED').first()
            if mr:
                initial.update({'maintenance_request': mr, 'asset': mr.asset, 'work_centre': mr.work_centre, 'priority': mr.priority, 'failure_mode': mr.failure_mode})
        return initial
    def form_valid(self, form):
        form.instance.created_by = self.request.user
        form.instance.plant = form.instance.asset.plant
        form.instance.status = 'DRAFT'
        if form.instance.maintenance_request_id and form.instance.maintenance_request.status not in ['ACCEPTED', 'WO_CREATED']:
            form.add_error('maintenance_request', 'The maintenance request must be accepted during screening before a work order can be created.')
            return self.form_invalid(form)
        response = super().form_valid(form)
        if form.instance.maintenance_request_id:
            MaintenanceRequest.objects.filter(pk=form.instance.maintenance_request_id).update(status='WO_CREATED')
        return response


class WorkOrderUpdateView(SecureUpdateView):
    model = WorkOrder; form_class = WorkOrderForm; title = 'Update Work Order Planning / Execution Data'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        obj = self.get_object()
        if obj.technical_lock:
            messages.error(request, 'This work order is technically completed (TECO) and technical fields are locked.')
            return redirect('work_order_detail', pk=obj.pk)
        return super().dispatch(request, *args, **kwargs)
    def form_valid(self, form):
        form.instance.created_by = self.object.created_by
        form.instance.plant = form.instance.asset.plant
        form.instance.status = self.object.status
        return super().form_valid(form)


def maintenance_user_can_control_work_order(user, work_order):
    """Record-level execution guard: maintenance engineers control only work assigned to them."""
    role = user_role(user)
    if role in {'Admin', 'HOD', 'Manager'}:
        return True
    if role != 'Maintenance':
        return False
    if work_order.assigned_to_id != user.id:
        return work_order.operations.filter(assigned_to_id=user.id).exists()
    return True


@login_required
@require_POST
def work_order_action(request, pk, action):
    if not role_can_access_url(request.user, 'work_order_action'):
        raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=pk)
    role = user_role(request.user)
    if role == 'Maintenance' and obj.assigned_to_id != request.user.id and not maintenance_user_can_control_work_order(request.user, obj):
        raise PermissionDenied('Maintenance engineers may only control work orders assigned to them or their operations.')
    if action in {'approve', 'verify', 'teco'} and role not in {'Admin', 'HOD', 'Manager'}:
        raise PermissionDenied
    if action == 'cost_close' and role not in {'Admin', 'HOD', 'Manager', 'Accounts'}:
        raise PermissionDenied
    if action == 'close' and role not in {'Admin', 'HOD', 'Manager'}:
        raise PermissionDenied
    if action in {'plan', 'submit_approval', 'prepare', 'ready_schedule', 'schedule', 'dispatch', 'release', 'start', 'complete'} and role not in {'Admin', 'Maintenance', 'HOD', 'Manager'}:
        raise PermissionDenied
    try:
        previous = serialise_instance(obj)
        transition_work_order(obj, action, request.user, request.POST.get('notes', ''))
        log_action(request, f'WO_{action.upper()}', obj, previous=previous)
        messages.success(request, f'{obj.wo_number} moved to {obj.get_status_display()}.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('work_order_detail', pk=obj.pk)


class WorkOrderOperationCreateView(SecureCreateView):
    model = WorkOrderOperation; form_class = WorkOrderOperationForm; title = 'Add Work Order Operation'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        self.work_order = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=kwargs['pk'])
        if self.work_order.technical_lock:
            messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=self.work_order.pk)
        return super().dispatch(request, *args, **kwargs)
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.request.method in {'POST', 'PUT'}:
            kwargs['instance'] = WorkOrderOperation(work_order=self.work_order)
        return kwargs
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        qs = self.work_order.operations.all()
        if 'parent_operation' in form.fields: form.fields['parent_operation'].queryset = qs
        if 'predecessor' in form.fields: form.fields['predecessor'].queryset = qs
        return form
    def form_valid(self, form):
        form.instance.work_order = self.work_order
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.work_order.pk})


class WorkOrderOperationUpdateView(SecureUpdateView):
    model = WorkOrderOperation; form_class = WorkOrderOperationForm; title = 'Update Work Order Operation'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        obj = self.get_object()
        if obj.work_order.technical_lock:
            messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=obj.work_order_id)
        return super().dispatch(request, *args, **kwargs)
    def get_queryset(self):
        qs = super().get_queryset().select_related('work_order')
        allowed_wo = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user).values_list('pk', flat=True)
        return qs.filter(work_order_id__in=allowed_wo)
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        qs = self.object.work_order.operations.exclude(pk=self.object.pk)
        if 'parent_operation' in form.fields: form.fields['parent_operation'].queryset = qs
        if 'predecessor' in form.fields: form.fields['predecessor'].queryset = qs
        return form
    def form_valid(self, form):
        form.instance.work_order = self.object.work_order
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.object.work_order_id})


class WorkOrderSpareCreateView(SecureCreateView):
    model = WorkOrderSpare; form_class = WorkOrderSpareForm; title = 'Plan Work Order Material'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        self.work_order = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=kwargs['pk'])
        if self.work_order.technical_lock:
            messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=self.work_order.pk)
        return super().dispatch(request, *args, **kwargs)
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.request.method in {'POST', 'PUT'}:
            kwargs['instance'] = WorkOrderSpare(work_order=self.work_order)
        return kwargs
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if 'operation' in form.fields: form.fields['operation'].queryset = self.work_order.operations.all()
        return form
    def form_valid(self, form):
        form.instance.work_order = self.work_order
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.work_order.pk})


class WorkOrderSpareUpdateView(SecureUpdateView):
    model = WorkOrderSpare; form_class = WorkOrderSpareForm; title = 'Update Planned Material'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        obj = self.get_object()
        if obj.work_order.technical_lock:
            messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=obj.work_order_id)
        return super().dispatch(request, *args, **kwargs)
    def get_queryset(self):
        qs = super().get_queryset().select_related('work_order')
        allowed_wo = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user).values_list('pk', flat=True)
        return qs.filter(work_order_id__in=allowed_wo)
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if 'operation' in form.fields:
            form.fields['operation'].queryset = self.object.work_order.operations.all()
        return form
    def form_valid(self, form):
        form.instance.work_order = self.object.work_order
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.object.work_order_id})


class WorkOrderLabourCreateView(SecureCreateView):
    model = WorkOrderLabour; form_class = WorkOrderLabourForm; title = 'Record Work Order Labour'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        self.work_order = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=kwargs['pk'])
        if self.work_order.technical_lock:
            messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=self.work_order.pk)
        return super().dispatch(request, *args, **kwargs)
    def form_valid(self, form):
        form.instance.work_order = self.work_order
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.work_order.pk})


class WorkOrderConfirmationCreateView(SecureCreateView):
    model = WorkOrderConfirmation; form_class = WorkOrderConfirmationForm; title = 'Operation Confirmation'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        self.work_order = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=kwargs['pk'])
        if self.work_order.technical_lock:
            messages.error(request, 'TECO work orders cannot receive new confirmations.')
            return redirect('work_order_detail', pk=self.work_order.pk)
        return super().dispatch(request, *args, **kwargs)
    def get_initial(self):
        initial = super().get_initial(); initial['person'] = self.request.user; return initial
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.request.method in {'POST', 'PUT'}:
            kwargs['instance'] = WorkOrderConfirmation(work_order=self.work_order)
        return kwargs
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        form.fields['operation'].queryset = self.work_order.operations.exclude(status='CANCELLED')
        return form
    def form_valid(self, form):
        if form.instance.operation.work_order_id != self.work_order.pk:
            form.add_error('operation', 'Select an operation from this work order.'); return self.form_invalid(form)
        if user_role(self.request.user) == 'Maintenance' and form.instance.operation.assigned_to_id != self.request.user.id:
            raise PermissionDenied('Maintenance engineers may only confirm operations assigned to them.')
        form.instance.work_order = self.work_order
        form.instance.person = self.request.user
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.work_order.pk})


@login_required
@require_POST
def work_order_operation_action(request, pk, action):
    if not role_can_access_url(request.user, 'work_order_operation_action'):
        raise PermissionDenied
    op = get_object_or_404(WorkOrderOperation.objects.select_related('work_order'), pk=pk)
    assert_object_access(request.user, op.work_order)
    if user_role(request.user) == 'Maintenance' and op.assigned_to_id != request.user.id:
        raise PermissionDenied('Maintenance engineers may only control operations assigned to them.')
    if action in {'start', 'resume', 'finish'} and op.predecessor_id and op.predecessor.status != 'COMPLETED':
        messages.error(request, f'Complete predecessor operation {op.predecessor.sequence} before this operation.')
        return redirect('work_order_detail', pk=op.work_order_id)
    if op.work_order.technical_lock:
        messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=op.work_order_id)
    if op.work_order.status != 'IN_PROGRESS':
        messages.error(request, 'Start the released work order before controlling operation execution.')
        return redirect('work_order_detail', pk=op.work_order_id)
    mapping = {
        'start': ({'READY','DISPATCHED','PAUSED'}, 'IN_PROGRESS'),
        'pause': ({'IN_PROGRESS'}, 'PAUSED'),
        'resume': ({'PAUSED'}, 'IN_PROGRESS'),
        'finish': ({'IN_PROGRESS'}, 'WORK_FINISHED'),
    }
    if action not in mapping:
        raise Http404
    allowed, target = mapping[action]
    if op.status not in allowed:
        messages.error(request, f'Operation cannot {action} from {op.get_status_display()}.')
        return redirect('work_order_detail', pk=op.work_order_id)
    from .models import OperationTimeEvent
    OperationTimeEvent.objects.create(operation=op, action={'start':'START','pause':'PAUSE','resume':'RESUME','finish':'STOP'}[action], performed_by=request.user)
    op.status = target; op.save(update_fields=['status','updated_at'])
    messages.success(request, f'Operation {op.sequence} moved to {op.get_status_display()}.')
    return redirect('work_order_detail', pk=op.work_order_id)


class SettlementRuleCreateView(SecureCreateView):
    model = SettlementRule; form_class = SettlementRuleForm; title = 'Add Settlement Rule'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        self.work_order = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=kwargs['pk'])
        return super().dispatch(request, *args, **kwargs)
    def form_valid(self, form):
        form.instance.work_order = self.work_order
        current = self.work_order.settlement_rules.aggregate(total=Sum('percentage'))['total'] or Decimal('0')
        if current + form.instance.percentage > Decimal('100'):
            form.add_error('percentage', 'Settlement percentages cannot exceed 100%.'); return self.form_invalid(form)
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.work_order.pk})


class WorkOrderRiskCreateView(SecureCreateView):
    model = RiskAssessment; form_class = RiskAssessmentForm; title = 'Add Work Order Risk Assessment'; success_url_name = 'work_order_list'
    def dispatch(self, request, *args, **kwargs):
        self.work_order = get_object_or_404(scope_queryset_for_user(WorkOrder.objects.all(), request.user), pk=kwargs['pk'])
        if self.work_order.technical_lock:
            messages.error(request, 'TECO work orders are technically locked.'); return redirect('work_order_detail', pk=self.work_order.pk)
        return super().dispatch(request, *args, **kwargs)
    def form_valid(self, form):
        form.instance.work_order = self.work_order
        return super().form_valid(form)
    def get_success_url(self): return reverse_lazy('work_order_detail', kwargs={'pk': self.work_order.pk})


# -----------------------------------------------------------------------------
# Maintenance History & Reliability - Excel replacement reporting layer
# -----------------------------------------------------------------------------

def _parse_report_period(request):
    today = timezone.localdate()
    default_start = today.replace(day=1)
    try:
        start_date = datetime.strptime(request.GET.get('from_date', ''), '%Y-%m-%d').date() if request.GET.get('from_date') else default_start
    except ValueError:
        start_date = default_start
    try:
        end_date = datetime.strptime(request.GET.get('to_date', ''), '%Y-%m-%d').date() if request.GET.get('to_date') else today
    except ValueError:
        end_date = today
    if end_date < start_date:
        start_date, end_date = end_date, start_date
    tz = timezone.get_current_timezone()
    period_start = timezone.make_aware(datetime.combine(start_date, time.min), tz)
    period_end = timezone.make_aware(datetime.combine(end_date + timedelta(days=1), time.min), tz)
    return start_date, end_date, period_start, period_end


def _maintenance_report_scope(request):
    start_date, end_date, period_start, period_end = _parse_report_period(request)
    assets = scope_queryset_for_user(Asset.objects.all(), request.user)
    work_orders = scope_queryset_for_user(WorkOrder.objects.all(), request.user).select_related(
        'asset', 'asset__functional_location', 'plant', 'work_centre', 'assigned_to', 'supervisor', 'verified_by'
    )
    downtime = scope_queryset_for_user(AssetDowntimeEvent.objects.all(), request.user)

    plant_id = request.GET.get('plant', '').strip()
    work_centre_id = request.GET.get('work_centre', '').strip()
    asset_id = request.GET.get('asset', '').strip()
    criticality = request.GET.get('criticality', '').strip()
    work_type = request.GET.get('work_type', '').strip()
    activity_type = request.GET.get('activity_type', '').strip()
    permit_type = request.GET.get('permit_type', '').strip()
    shift = request.GET.get('shift', '').strip()

    if plant_id:
        assets = assets.filter(plant_id=plant_id)
        work_orders = work_orders.filter(plant_id=plant_id)
        downtime = downtime.filter(asset__plant_id=plant_id)
    if asset_id:
        assets = assets.filter(pk=asset_id)
        work_orders = work_orders.filter(asset_id=asset_id)
        downtime = downtime.filter(asset_id=asset_id)
    if criticality:
        assets = assets.filter(criticality=criticality)
        work_orders = work_orders.filter(asset__criticality=criticality)
        downtime = downtime.filter(asset__criticality=criticality)
    if work_centre_id:
        work_orders = work_orders.filter(work_centre_id=work_centre_id)
        scoped_asset_ids = work_orders.values_list('asset_id', flat=True).distinct()
        assets = assets.filter(id__in=scoped_asset_ids)
        downtime = downtime.filter(asset_id__in=scoped_asset_ids)
    if work_type:
        work_orders = work_orders.filter(work_type=work_type)
    if activity_type:
        work_orders = work_orders.filter(activity_type=activity_type)
    if shift:
        work_orders = work_orders.filter(shift=shift)
    if permit_type:
        work_orders = work_orders.filter(permits__permit_type=permit_type).distinct()

    # Reliability calculations only need filtered WorkOrder rows, not the
    # select_related() joins used by the history UI. Keeping a lean queryset here
    # prevents conflicts if the reliability service applies only()/defer().
    reliability_orders = work_orders.select_related(None)
    history = work_orders.filter(
        Q(actual_end__gte=period_start, actual_end__lt=period_end)
        | Q(actual_end__isnull=True, created_at__gte=period_start, created_at__lt=period_end)
    ).distinct().order_by('-actual_end', '-created_at')

    return {
        'start_date': start_date, 'end_date': end_date, 'period_start': period_start, 'period_end': period_end,
        'assets': assets.distinct(), 'work_orders': reliability_orders.distinct(), 'history': history,
        'downtime': downtime.distinct(),
    }


class MaintenanceHistoryReliabilityView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/maintenance_history_reliability.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        scope = _maintenance_report_scope(self.request)
        metrics = calculate_reliability_metrics(
            assets=scope['assets'], work_orders=scope['work_orders'], downtime_events=scope['downtime'],
            period_start=scope['period_start'], period_end=scope['period_end'],
        )
        context.update(scope)
        context.update({
            'metrics': metrics,
            'plants': accessible_plants(self.request.user),
            'work_centres': scope_queryset_for_user(WorkCentre.objects.filter(active=True), self.request.user),
            'asset_options': scope_queryset_for_user(Asset.objects.filter(is_active=True), self.request.user).order_by('asset_id'),
            'criticality_choices': Asset.CRITICALITY_CHOICES,
            'work_type_choices': WorkOrder.WORK_TYPES,
            'activity_type_choices': WorkOrder.ACTIVITY_TYPES,
            'shift_choices': WorkOrder.SHIFT_CHOICES,
            'permit_type_choices': PermitToWork.PERMIT_TYPES,
        })
        return context


def _history_rows(request):
    scope = _maintenance_report_scope(request)
    rows = []
    for wo in scope['history'].prefetch_related('material_issues__item', 'permits'):
        permits = ', '.join(f'{p.permit_no} / {p.get_permit_type_display()}' for p in wo.permits.all()) or '-'
        spares = ', '.join(f'{i.item.item_no} x {i.quantity_issued}' for i in wo.material_issues.all() if i.status == 'POSTED') or '-'
        rows.append([
            (wo.actual_end or wo.created_at).date(), wo.wo_number, permits, wo.asset.tag_number or wo.asset.asset_id,
            wo.asset.name, wo.work_centre.code, wo.asset.get_criticality_display(), wo.get_work_type_display(),
            wo.get_activity_type_display(), spares, wo.get_shift_display(), wo.repair_hours or '',
            wo.assigned_to.get_full_name() or wo.assigned_to.username if wo.assigned_to else '-',
            wo.supervisor.get_full_name() or wo.supervisor.username if wo.supervisor else '-', wo.actual_cost, wo.completion_notes or '-',
        ])
    return scope, rows


@login_required
def maintenance_history_excel(request):
    if not role_can_access_url(request.user, 'maintenance_history_excel'):
        raise PermissionDenied
    from openpyxl import Workbook
    from openpyxl.styles import Font
    scope, rows = _history_rows(request)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Maintenance History'
    headers = ['Date', 'WO', 'Permit No. / Type', 'TAG No.', 'Equipment', 'Work Centre', 'Criticality', 'Work Type', 'Activity', 'Spares Consumed', 'Shift', 'Repair Hours', 'Work Done By', 'Supervisor', 'Actual Cost', 'Remarks']
    ws.append(headers)
    for cell in ws[1]: cell.font = Font(bold=True)
    for row in rows: ws.append(row)
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        width = min(max(len(str(c.value or '')) for c in col) + 2, 45)
        ws.column_dimensions[col[0].column_letter].width = width
    metrics = calculate_reliability_metrics(scope['assets'], scope['work_orders'], scope['downtime'], scope['period_start'], scope['period_end'])
    kpi = wb.create_sheet('Reliability KPIs')
    kpi.append(['KPI', 'Value'])
    kpi.append(['Period', f"{scope['start_date']} to {scope['end_date']}"])
    kpi.append(['Breakdowns', metrics['breakdown_failures']])
    kpi.append(['Breakdown Hours', metrics['breakdown_downtime_hours']])
    kpi.append(['MTTR Hours', metrics['mttr_hours'] if metrics['mttr_hours'] is not None else 'N/A'])
    kpi.append(['MTBF Hours', metrics['mtbf_hours'] if metrics['mtbf_hours'] is not None else 'N/A'])
    kpi.append(['Availability %', metrics['availability_percent'] if metrics['availability_percent'] is not None else 'N/A'])
    kpi.append(['Failure Rate / 1000 h', metrics['failure_rate_per_1000_hours'] if metrics['failure_rate_per_1000_hours'] is not None else 'N/A'])
    for cell in kpi[1]: cell.font = Font(bold=True)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="maintenance-history-reliability.xlsx"'
    wb.save(response)
    return response


@login_required
def maintenance_history_pdf(request):
    if not role_can_access_url(request.user, 'maintenance_history_pdf'):
        raise PermissionDenied
    scope, rows = _history_rows(request)
    metrics = calculate_reliability_metrics(scope['assets'], scope['work_orders'], scope['downtime'], scope['period_start'], scope['period_end'])
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="maintenance-history-reliability.pdf"'
        doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=18, leftMargin=18, topMargin=20, bottomMargin=20)
        styles = getSampleStyleSheet()
        story = [Paragraph('Maintenance History & Reliability', styles['Title'])]
        story.append(Paragraph(
            f"Period: {scope['start_date']} to {scope['end_date']} | Breakdowns: {metrics['breakdown_failures']} | MTTR: {metrics['mttr_hours'] if metrics['mttr_hours'] is not None else '-'} h | MTBF: {metrics['mtbf_hours'] if metrics['mtbf_hours'] is not None else '-'} h | Availability: {metrics['availability_percent'] if metrics['availability_percent'] is not None else '-'}% | Failure Rate: {metrics['failure_rate_per_1000_hours'] if metrics['failure_rate_per_1000_hours'] is not None else '-'} /1000h",
            styles['Normal']))
        story.append(Spacer(1, 10))
        headers = ['Date', 'WO', 'Permit', 'TAG', 'Equipment', 'WC', 'Crit.', 'Work Type', 'Activity', 'Spares', 'Shift', 'Repair h', 'Technician', 'Supervisor', 'Cost', 'Remarks']
        data = [headers] + [[str(v)[:32] for v in row] for row in rows]
        table = Table(data, repeatRows=1, colWidths=[42,52,65,50,65,32,36,55,50,60,32,38,55,55,45,75])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.lightgrey), ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 6), ('GRID', (0,0), (-1,-1), 0.25, colors.grey),
            ('VALIGN', (0,0), (-1,-1), 'TOP'), ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.whitesmoke]),
        ]))
        story.append(table)
        doc.build(story)
        return response
    except ImportError:
        return HttpResponse('ReportLab is required for PDF export.', status=500)


class DowntimeListView(ModuleListView):
    model = AssetDowntimeEvent; title = 'Equipment Downtime / Availability'
    columns = [('asset.asset_id', 'Asset'), ('downtime_type', 'Type'), ('started_at', 'Start'), ('ended_at', 'End'), ('duration_hours', 'Hours'), ('reason', 'Reason')]
    create_url_name = 'downtime_create'
    def get_select_related(self): return ['asset', 'work_order', 'recorded_by']


class DowntimeCreateView(SecureCreateView):
    model = AssetDowntimeEvent; form_class = DowntimeForm; title = 'Record Equipment Downtime'; success_url_name = 'downtime_list'
    def get_initial(self):
        initial = super().get_initial()
        wo_id = self.request.GET.get('work_order')
        if wo_id:
            wo = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user).filter(pk=wo_id).first()
            if wo:
                initial.update({'work_order': wo, 'asset': wo.asset, 'downtime_type': 'BREAKDOWN' if wo.work_type == 'BREAKDOWN' else 'PLANNED'})
        return initial
    def form_valid(self, form):
        form.instance.recorded_by = self.request.user
        return super().form_valid(form)


class PMPlanListView(ModuleListView):
    model = PMPlan; title = 'Preventive Maintenance Plans'
    columns = [('name', 'Plan'), ('asset.asset_id', 'Asset'), ('work_centre.code', 'Work Centre'), ('frequency_value', 'Frequency'), ('frequency_unit', 'Unit'), ('next_due_date', 'Next Due'), ('active', 'Active')]
    create_url_name = 'pm_plan_create'; detail_url_name = 'pm_plan_detail'; update_url_name = 'pm_plan_update'
    def get_select_related(self): return ['asset', 'plant', 'work_centre', 'responsible_user']


class PMPlanDetailView(RoleRequiredMixin, DetailView):
    model = PMPlan
    template_name = 'asset_mgmt/pm_plan_detail.html'
    context_object_name = 'plan'
    def get_queryset(self):
        return scope_queryset_for_user(PMPlan.objects.select_related('asset','plant','work_centre','responsible_user','task_list','strategy','counter_meter'), self.request.user)
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        due, reasons = pm_plan_is_due(self.object)
        context.update({
            'is_due': due, 'due_reasons': reasons,
            'counters': self.object.counters.select_related('meter'),
            'strategy_calls': self.object.strategy_packages.select_related('package'),
            'legacy_tasks': self.object.tasks.all(),
        })
        return context


class PMPlanCreateView(SecureCreateView):
    model = PMPlan; form_class = PMPlanForm; title = 'Create PM Plan'; success_url_name = 'pm_plan_list'
    def form_valid(self, form):
        form.instance.plant = form.instance.asset.plant
        return super().form_valid(form)


class PMPlanUpdateView(SecureUpdateView):
    model = PMPlan; form_class = PMPlanForm; title = 'Update PM Plan'; success_url_name = 'pm_plan_list'
    def form_valid(self, form):
        form.instance.plant = form.instance.asset.plant
        return super().form_valid(form)


class PMTaskCreateView(SecureCreateView):
    model = PMTask; form_class = PMTaskForm; title = 'Add PM Task'; success_url_name = 'pm_plan_list'


@login_required
@require_POST
def pm_generate_wo(request, pk):
    if not role_can_access_url(request.user, 'pm_generate_wo'): raise PermissionDenied
    plan = get_object_or_404(scope_queryset_for_user(PMPlan.objects.select_related('asset', 'work_centre'), request.user), pk=pk)
    try:
        wo = generate_pm_work_order(plan, request.user)
        log_action(request, 'CREATE', wo, remarks=f'Generated from PM plan {plan.pk}')
        messages.success(request, f'Work order {wo.wo_number} generated from PM plan.')
        return redirect('work_order_detail', pk=wo.pk)
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
        return redirect('pm_plan_list')


class PMPlanCounterCreateView(SecureCreateView):
    model = PMPlanCounter; form_class = PMPlanCounterForm; title = 'Add PM Counter'; success_url_name = 'pm_plan_list'
    def get_initial(self):
        initial = super().get_initial(); initial['pm_plan'] = self.request.GET.get('pm_plan') or None; return initial
    def get_success_url(self):
        return reverse_lazy('pm_plan_detail', kwargs={'pk': self.object.pm_plan_id})


class PMPlanStrategyPackageCreateView(SecureCreateView):
    model = PMPlanStrategyPackage; form_class = PMPlanStrategyPackageForm; title = 'Add PM Strategy Package Call'; success_url_name = 'pm_plan_list'
    def get_initial(self):
        initial = super().get_initial(); initial['pm_plan'] = self.request.GET.get('pm_plan') or None; return initial
    def get_success_url(self):
        return reverse_lazy('pm_plan_detail', kwargs={'pk': self.object.pm_plan_id})


class MaintenanceTaskListListView(ModuleListView):
    model = MaintenanceTaskList; title = 'Maintenance Task Lists'
    columns = [('code','Code'),('name','Name'),('plant.code','Plant'),('list_type','Type'),('revision','Revision'),('active','Active')]
    create_url_name = 'task_list_create'; detail_url_name = 'task_list_detail'
    def get_select_related(self): return ['plant','asset','functional_location','strategy']


class MaintenanceTaskListDetailView(RoleRequiredMixin, DetailView):
    model = MaintenanceTaskList
    template_name = 'asset_mgmt/task_list_detail.html'
    context_object_name = 'task_list'
    def get_queryset(self):
        return scope_queryset_for_user(MaintenanceTaskList.objects.select_related('plant','asset','functional_location','strategy'), self.request.user)
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['operations'] = self.object.operations.select_related('work_centre','strategy_package','service_item').prefetch_related('materials__spare_part')
        return context


class MaintenanceTaskListCreateView(SecureCreateView):
    model = MaintenanceTaskList; form_class = MaintenanceTaskListForm; title = 'Create Maintenance Task List'; success_url_name = 'task_list_list'


class MaintenanceTaskListOperationCreateView(SecureCreateView):
    model = MaintenanceTaskListOperation; form_class = MaintenanceTaskListOperationForm; title = 'Add Task List Operation'; success_url_name = 'task_list_list'
    def get_initial(self):
        initial = super().get_initial(); initial['task_list'] = self.request.GET.get('task_list') or None; return initial
    def get_success_url(self): return reverse_lazy('task_list_detail', kwargs={'pk': self.object.task_list_id})


class MaintenanceTaskListMaterialCreateView(SecureCreateView):
    model = MaintenanceTaskListMaterial; form_class = MaintenanceTaskListMaterialForm; title = 'Add Task List Material'; success_url_name = 'task_list_list'
    def get_initial(self):
        initial = super().get_initial(); initial['task_operation'] = self.request.GET.get('operation') or None; return initial
    def get_success_url(self): return reverse_lazy('task_list_detail', kwargs={'pk': self.object.task_operation.task_list_id})


class MaintenanceStrategyListView(ModuleListView):
    model = MaintenanceStrategy; title = 'Maintenance Strategies'
    columns = [('plant.code','Plant'),('name','Strategy'),('description','Description'),('active','Active')]
    create_url_name = 'maintenance_strategy_create'
    def get_select_related(self): return ['plant']


class MaintenanceStrategyCreateView(SecureCreateView):
    model = MaintenanceStrategy; form_class = MaintenanceStrategyForm; title = 'Create Maintenance Strategy'; success_url_name = 'maintenance_strategy_list'


class MaintenanceStrategyPackageCreateView(SecureCreateView):
    model = MaintenanceStrategyPackage; form_class = MaintenanceStrategyPackageForm; title = 'Add Strategy Package'; success_url_name = 'maintenance_strategy_list'


class ConditionRuleListView(ModuleListView):
    model = ConditionRule; title = 'Condition-based Maintenance Rules'
    columns = [('meter.asset.asset_id','Asset'),('meter.meter_type','Meter'),('name','Rule'),('operator','Operator'),('threshold','Threshold'),('priority','Priority'),('active','Active')]
    create_url_name = 'condition_rule_create'
    def get_select_related(self): return ['meter','meter__asset','work_centre']


class ConditionRuleCreateView(SecureCreateView):
    model = ConditionRule; form_class = ConditionRuleForm; title = 'Create Condition Rule'; success_url_name = 'condition_rule_list'


class SchedulingBoardView(RoleRequiredMixin, TemplateView):
    template_name = 'asset_mgmt/scheduling_board.html'
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plant, plants = selected_plant(self.request)
        today = timezone.localdate()
        try:
            horizon = max(1, min(31, int(self.request.GET.get('days','7'))))
        except ValueError:
            horizon = 7
        end_date = today + timedelta(days=horizon-1)
        centres = scope_queryset_for_user(WorkCentre.objects.filter(active=True), self.request.user)
        orders = scope_queryset_for_user(WorkOrder.objects.filter(status__in=['READY_TO_SCHEDULE','SCHEDULED','DISPATCHED']), self.request.user).select_related('asset','work_centre','assigned_to')
        if plant:
            centres = centres.filter(plant=plant); orders = orders.filter(plant=plant)
        rows = []
        for wc in centres:
            rows.append({'work_centre':wc, **work_centre_capacity(wc,today,end_date)})
        context.update({'plants':plants,'selected_plant':plant,'start_date':today,'end_date':end_date,'horizon':horizon,'capacity_rows':rows,'schedule_orders':orders.order_by('planned_start')})
        return context


class WorkCentreCapacityUpdateView(SecureUpdateView):
    model = WorkCentre; form_class = WorkCentreCapacityForm; title = 'Work Centre Capacity'; success_url_name = 'scheduling_board'


class ShutdownListView(ModuleListView):
    model = ShutdownPlan; title = 'Shutdown Maintenance Plans'
    columns = [('plan_no', 'Plan'), ('plant.code', 'Plant'), ('title', 'Title'), ('start_date', 'Start'), ('end_date', 'End'), ('completion_percent', 'Progress %'), ('status', 'Status')]
    create_url_name = 'shutdown_create'; update_url_name = 'shutdown_update'
    def get_select_related(self): return ['plant', 'coordinator']

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        page_items = context.get('items')
        plans = list(page_items.object_list if hasattr(page_items, 'object_list') else page_items or [])
        context['shutdown_chart_json'] = chart_json({
            'labels': [plan.plan_no for plan in plans],
            'values': [plan.completion_percent for plan in plans],
        })
        return context


class ShutdownCreateView(SecureCreateView):
    model = ShutdownPlan; form_class = ShutdownPlanForm; title = 'Create Shutdown Plan'; success_url_name = 'shutdown_list'


class ShutdownUpdateView(SecureUpdateView):
    model = ShutdownPlan; form_class = ShutdownPlanForm; title = 'Update Shutdown Plan'; success_url_name = 'shutdown_list'


class ShutdownJobCreateView(SecureCreateView):
    model = ShutdownJob; form_class = ShutdownJobForm; title = 'Add Shutdown Job'; success_url_name = 'shutdown_list'


class RCAListView(ModuleListView):
    model = RCARecord; title = 'Root Cause Analysis'
    columns = [('rca_no', 'RCA'), ('work_order.wo_number', 'Work Order'), ('method', 'Method'), ('failure_mode', 'Failure Mode'), ('responsible_person.username', 'Owner'), ('target_date', 'Target'), ('closed', 'Closed')]
    create_url_name = 'rca_create'
    def get_select_related(self): return ['work_order', 'responsible_person']


class RCACreateView(SecureCreateView):
    model = RCARecord; form_class = RCAForm; title = 'Create RCA (5-Why / Fishbone)'; success_url_name = 'rca_list'


class FishboneCreateView(SecureCreateView):
    model = FishboneCause; form_class = FishboneCauseForm; title = 'Add Fishbone Cause'; success_url_name = 'rca_list'


class CAPAListView(ModuleListView):
    model = CAPAAction; title = 'Corrective and Preventive Actions'
    columns = [('rca.rca_no', 'RCA'), ('action_type', 'Type'), ('description', 'Action'), ('owner.username', 'Owner'), ('due_date', 'Due'), ('score', 'Score'), ('status', 'Status')]
    create_url_name = 'capa_create'
    def get_select_related(self): return ['rca', 'rca__work_order', 'owner']


class CAPACreateView(SecureCreateView):
    model = CAPAAction; form_class = CAPAForm; title = 'Create CAPA Action'; success_url_name = 'capa_list'


class LLFListView(ModuleListView):
    model = LLFObservation; title = 'LLF / TPM Observations'
    columns = [('observation_no', 'Observation'), ('asset.asset_id', 'Asset'), ('sense', 'Sense'), ('severity', 'Severity'), ('assigned_work_centre.code', 'Work Centre'), ('status', 'Status'), ('observed_at', 'Observed At')]
    create_url_name = 'llf_create'
    def get_select_related(self): return ['asset', 'plant', 'observed_by', 'assigned_work_centre']


class LLFCreateView(SecureCreateView):
    model = LLFObservation; form_class = LLFForm; title = 'Create LLF Observation'; success_url_name = 'llf_list'
    def form_valid(self, form):
        form.instance.observed_by = self.request.user
        form.instance.plant = form.instance.asset.plant
        return super().form_valid(form)


@login_required
@require_POST
def llf_convert_mr(request, pk):
    if not role_can_access_url(request.user, 'llf_convert_mr'): raise PermissionDenied
    observation = get_object_or_404(scope_queryset_for_user(LLFObservation.objects.all(), request.user), pk=pk)
    if observation.maintenance_request_id:
        messages.info(request, 'A maintenance request already exists for this observation.')
        return redirect('llf_list')
    mr = MaintenanceRequest.objects.create(
        plant=observation.plant, asset=observation.asset, work_centre=observation.assigned_work_centre,
        reported_by=request.user, reported_department='MAINTENANCE', problem_description=observation.observation,
        priority='URGENT' if observation.severity == 'CRITICAL' else ('HIGH' if observation.severity == 'HIGH' else 'MEDIUM'),
        failure_mode='LLF / TPM observation',
    )
    observation.maintenance_request = mr; observation.status = 'MR_CREATED'; observation.save(update_fields=['maintenance_request', 'status', 'updated_at'])
    log_action(request, 'CREATE', mr, remarks=f'Converted from {observation.observation_no}')
    messages.success(request, f'Maintenance request {mr.request_no} created.')
    return redirect('maintenance_request_list')


class MOCListView(ModuleListView):
    model = ModificationOrder; title = 'Management of Change (MOC)'
    columns = [('mod_order_no', 'MOC'), ('title', 'Title'), ('asset.asset_id', 'Asset'), ('work_centre.code', 'Work Centre'), ('assigned_approver.username', 'Approver'), ('status', 'Status'), ('requested_date', 'Date')]
    create_url_name = 'moc_create'; update_url_name = 'moc_update'
    def get_select_related(self): return ['asset', 'plant', 'work_centre', 'assigned_approver']


class MOCCreateView(SecureCreateView):
    model = ModificationOrder; form_class = ModificationOrderForm; title = 'Create Management of Change'; success_url_name = 'moc_list'
    def form_valid(self, form):
        form.instance.requested_by = self.request.user
        form.instance.plant = form.instance.asset.plant
        return super().form_valid(form)


class MOCUpdateView(SecureUpdateView):
    model = ModificationOrder; form_class = ModificationOrderForm; title = 'Update Management of Change'; success_url_name = 'moc_list'
    def dispatch(self, request, *args, **kwargs):
        obj = self.get_queryset().filter(pk=kwargs['pk']).first()
        if obj and obj.is_locked and not is_admin(request.user):
            messages.error(request, 'Approved MOC records are locked. Use the PDF copy for end-user reference.')
            return redirect('moc_list')
        return super().dispatch(request, *args, **kwargs)
    def form_valid(self, form):
        form.instance.requested_by = self.object.requested_by
        form.instance.plant = form.instance.asset.plant
        form.instance.status = self.object.status
        return super().form_valid(form)


@login_required
@require_POST
def moc_action(request, pk, action):
    if not role_can_access_url(request.user, 'moc_action'):
        raise PermissionDenied
    try:
        with transaction.atomic():
            queryset = scope_queryset_for_user(ModificationOrder.objects.select_for_update(), request.user)
            moc = get_object_or_404(queryset, pk=pk)
            previous = moc.status
            remarks = request.POST.get('remarks', '').strip()
            if action == 'submit' and moc.status in ['DRAFT', 'REJECTED'] and moc.requested_by_id == request.user.id:
                moc.status = 'SUBMITTED'
                moc.approval_remarks = ''
            elif action == 'approve' and moc.status in ['SUBMITTED', 'REVIEWED'] and (moc.assigned_approver_id == request.user.id or is_admin(request.user)):
                if moc.requested_by_id == request.user.id:
                    raise PermissionDenied('Self-approval is not allowed.')
                if len(remarks) < 3:
                    raise ValidationError('Approval remarks are required.')
                moc.status = 'APPROVED'
                moc.approved_by = request.user
                moc.approved_at = timezone.now()
                moc.locked_at = timezone.now()
                moc.approval_remarks = remarks
            elif action == 'reject' and moc.status in ['SUBMITTED', 'REVIEWED'] and (moc.assigned_approver_id == request.user.id or is_admin(request.user)):
                if len(remarks) < 5:
                    raise ValidationError('A rejection reason is required.')
                moc.status = 'REJECTED'
                moc.approval_remarks = remarks
            else:
                raise ValidationError('Invalid MOC status transition or permission.')
            moc.save()
            log_action(request, 'STATUS', moc, previous={'status': previous}, remarks=f'MOC action: {action}. {remarks}')
        messages.success(request, f'MOC {moc.mod_order_no} is now {moc.get_status_display()}.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('moc_list')


@login_required
def moc_pdf(request, pk):
    if not role_can_access_url(request.user, 'moc_pdf'): raise PermissionDenied
    moc = get_object_or_404(scope_queryset_for_user(ModificationOrder.objects.all(), request.user), pk=pk)
    lines = [('MOC No', moc.mod_order_no), ('Status', moc.get_status_display()), ('Plant', moc.plant), ('Asset', moc.asset), ('Title', moc.title), ('Current Condition', moc.current_condition), ('Proposed Change', moc.proposed_change), ('Reason', moc.reason), ('Safety Impact', moc.safety_impact), ('Quality Impact', moc.quality_impact), ('Environment Impact', moc.environment_impact), ('Production Impact', moc.production_impact), ('Risk Assessment', moc.risk_assessment), ('Implementation Plan', moc.implementation_plan), ('Rollback Plan', moc.rollback_plan), ('Validation Plan', moc.validation_plan), ('Requested By', moc.requested_by), ('Approved By', moc.approved_by), ('Approved At', moc.approved_at)]
    return safe_pdf_response('Management of Change Approval Copy', lines, f'{moc.mod_order_no}.pdf')


# Existing compliance modules retained
class InspectionPlanListView(ModuleListView):
    model = InspectionPlan; title = 'Inspection Plans'; columns = [('asset.asset_id', 'Asset'), ('inspection_type', 'Type'), ('statutory_form', 'Form'), ('next_due_date', 'Due'), ('active', 'Active')]; create_url_name = 'inspection_plan_create'
    def get_select_related(self): return ['asset', 'plant']
class InspectionPlanCreateView(SecureCreateView):
    model = InspectionPlan; form_class = InspectionPlanForm; title = 'Create Inspection Plan'; success_url_name = 'inspection_plan_list'
    def form_valid(self, form): form.instance.plant = form.instance.asset.plant; return super().form_valid(form)
class CalibrationPlanListView(ModuleListView):
    model = CalibrationPlan; title = 'Calibration Plans'; columns = [('asset.asset_id', 'Asset'), ('instrument_range', 'Range'), ('accuracy', 'Accuracy'), ('next_due_date', 'Due'), ('active', 'Active')]; create_url_name = 'calibration_plan_create'
    def get_select_related(self): return ['asset', 'plant']
class CalibrationPlanCreateView(SecureCreateView):
    model = CalibrationPlan; form_class = CalibrationPlanForm; title = 'Create Calibration Plan'; success_url_name = 'calibration_plan_list'
    def form_valid(self, form): form.instance.plant = form.instance.asset.plant; return super().form_valid(form)

class InspectionRecordCreateView(SecureCreateView):
    model = InspectionRecord; form_class = InspectionRecordForm; title = 'Record Inspection'; success_url_name = 'inspection_plan_list'
    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get('plan'): initial['plan'] = self.request.GET.get('plan')
        return initial


class CalibrationRecordCreateView(SecureCreateView):
    model = CalibrationRecord; form_class = CalibrationRecordForm; title = 'Record Calibration'; success_url_name = 'calibration_plan_list'
    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get('plan'): initial['plan'] = self.request.GET.get('plan')
        if not initial.get('calibrated_by'):
            initial['calibrated_by'] = self.request.user.get_full_name() or self.request.user.username
        return initial


class ShutdownJobUpdateView(SecureUpdateView):
    model = ShutdownJob; form_class = ShutdownJobForm; title = 'Update Shutdown Job'; success_url_name = 'shutdown_list'


class PermitListView(ModuleListView):
    model = PermitToWork; title = 'Permit to Work'; columns = [('permit_no', 'Permit'), ('work_order.wo_number', 'WO'), ('permit_type', 'Type'), ('valid_from', 'Valid From'), ('valid_to', 'Valid To'), ('status', 'Status')]; create_url_name = 'permit_create'
    def get_select_related(self): return ['work_order']
class PermitCreateView(SecureCreateView):
    model = PermitToWork; form_class = PermitToWorkForm; title = 'Create Permit to Work'; success_url_name = 'permit_list'
    def get_initial(self):
        initial = super().get_initial()
        wo_id = self.request.GET.get('work_order')
        if wo_id:
            wo = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user).filter(pk=wo_id).first()
            if wo:
                initial.update({'work_order': wo})
        return initial
    def form_valid(self, form):
        form.instance.status = 'REQUESTED'
        return super().form_valid(form)


@login_required
@require_POST
def permit_action(request, pk, action):
    if not role_can_access_url(request.user, 'permit_action'):
        raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(PermitToWork.objects.all(), request.user), pk=pk)
    role = user_role(request.user)
    if action in {'review', 'approve', 'reject', 'issue'} and role not in {'Admin', 'Safety', 'HOD'}:
        raise PermissionDenied
    if action == 'close' and role not in {'Admin', 'Safety', 'Maintenance', 'HOD'}:
        raise PermissionDenied
    try:
        previous = serialise_instance(obj)
        transition_permit(obj, action, request.user, request.POST.get('notes', ''))
        log_action(request, f'PERMIT_{action.upper()}', obj, previous=previous)
        messages.success(request, f'{obj.permit_no} moved to {obj.get_status_display()}.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('permit_list')


# -----------------------------------------------------------------------------
# Materials, one-stage PR approval and secure stock posting (13, 16, 17, 18, 24)
# -----------------------------------------------------------------------------

class SparePartListView(ModuleListView):
    model = SparePart; title = 'Engineering Store Item Master'
    columns = [('item_no', 'Item No'), ('name', 'Name'), ('plant.code', 'Plant'), ('item_description', 'Description'), ('current_stock', 'Quantity'), ('minimum_stock', 'Min'), ('item_rate', 'Rate'), ('stock_value', 'Value')]
    create_url_name = 'spare_create'; update_url_name = 'spare_update'
    def get_select_related(self): return ['plant']
    def get_queryset(self):
        qs = super().get_queryset()
        if self.request.GET.get('stock') == 'low': qs = qs.filter(current_stock__lte=F('minimum_stock'))
        return qs
class SparePartCreateView(SecureCreateView): model = SparePart; form_class = SparePartForm; title = 'Add Item'; success_url_name = 'spare_list'
class SparePartUpdateView(SecureUpdateView): model = SparePart; form_class = SparePartForm; title = 'Update Item'; success_url_name = 'spare_list'


class PurchaseRequestListView(ModuleListView):
    model = PurchaseRequest; title = 'Purchase Requests / Indents'
    columns = [('request_no', 'PR'), ('item.item_no', 'Item'), ('quantity', 'Qty'), ('plant.code', 'Plant'), ('work_centre.code', 'Work Centre'), ('requested_by.username', 'Requester'), ('assigned_hod.username', 'HOD'), ('status', 'Status')]
    create_url_name = 'purchase_request_create'; detail_url_name = 'purchase_request_detail'; update_url_name = 'purchase_request_update'
    def get_select_related(self): return ['item', 'plant', 'work_centre', 'requested_by', 'assigned_hod']


class PurchaseRequestDetailView(RoleRequiredMixin, PlantScopedMixin, DetailView):
    model = PurchaseRequest; template_name = 'asset_mgmt/purchase_request_detail.html'; context_object_name = 'purchase_request'
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        pr = self.object
        context['approval_form'] = ApprovalRemarksForm()
        context['rejection_form'] = RejectionForm()
        context['can_edit_pr'] = pr.status in ['DRAFT', 'REJECTED'] and (pr.requested_by_id == self.request.user.id or is_admin(self.request.user))
        context['can_submit_pr'] = context['can_edit_pr'] and role_can_access_url(self.request.user, 'purchase_request_submit')
        context['can_approve_pr'] = (
            pr.status == 'SUBMITTED'
            and role_can_access_url(self.request.user, 'purchase_request_approve')
            and (pr.assigned_hod_id == self.request.user.id or is_admin(self.request.user) or is_plant_head_for(self.request.user, pr.plant_id))
            and pr.requested_by_id != self.request.user.id
        )
        context['can_process_pr'] = pr.status == 'APPROVED' and role_can_access_url(self.request.user, 'purchase_request_processing')
        context['can_create_po'] = pr.status in ['APPROVED', 'PROCESSING'] and role_can_access_url(self.request.user, 'purchase_order_create') and not pr.purchase_orders.exists()
        return context


class PurchaseRequestCreateView(SecureCreateView):
    model = PurchaseRequest; form_class = PurchaseRequestForm; title = 'Raise Purchase Request'; success_url_name = 'purchase_request_list'
    def form_valid(self, form):
        form.instance.requested_by = self.request.user
        try: form.instance.assigned_hod = choose_hod(self.request.user, form.instance.plant, form.instance.work_centre)
        except ValidationError as exc:
            form.add_error('work_centre', exc); return self.form_invalid(form)
        response = super().form_valid(form)
        PurchaseRequestApprovalHistory.objects.create(purchase_request=self.object, action='CREATED', previous_status='', new_status='DRAFT', performed_by=self.request.user)
        return response


class PurchaseRequestUpdateView(SecureUpdateView):
    model = PurchaseRequest; form_class = PurchaseRequestForm; title = 'Update Purchase Request'; success_url_name = 'purchase_request_list'
    def dispatch(self, request, *args, **kwargs):
        obj = self.get_queryset().filter(pk=kwargs['pk']).first()
        if obj and obj.status not in ['DRAFT', 'REJECTED']:
            messages.error(request, 'Only Draft or Rejected purchase requests can be edited.')
            return redirect('purchase_request_detail', pk=obj.pk)
        if obj and obj.requested_by_id != request.user.id and not is_admin(request.user): raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)
    def form_valid(self, form):
        form.instance.requested_by = self.object.requested_by
        form.instance.status = self.object.status
        form.instance.assigned_hod = choose_hod(self.object.requested_by, form.instance.plant, form.instance.work_centre)
        return super().form_valid(form)


@login_required
@require_POST
def purchase_request_submit(request, pk):
    if not role_can_access_url(request.user, 'purchase_request_submit'): raise PermissionDenied
    pr = get_object_or_404(scope_queryset_for_user(PurchaseRequest.objects.all(), request.user), pk=pk)
    try: submit_purchase_request(pr, request.user); messages.success(request, 'Purchase request submitted to HOD.')
    except (ValidationError, PermissionDenied) as exc: messages.error(request, str(exc))
    return redirect('purchase_request_detail', pk=pk)


@login_required
@require_POST
def purchase_request_approve(request, pk):
    if not role_can_access_url(request.user, 'purchase_request_approve'): raise PermissionDenied
    pr = get_object_or_404(scope_queryset_for_user(PurchaseRequest.objects.all(), request.user), pk=pk)
    form = ApprovalRemarksForm(request.POST)
    if form.is_valid():
        try: approve_purchase_request(pr, request.user, form.cleaned_data['remarks']); messages.success(request, 'Purchase request approved.')
        except (ValidationError, PermissionDenied) as exc: messages.error(request, str(exc))
    else: messages.error(request, 'Approval remarks are required.')
    return redirect('purchase_request_detail', pk=pk)


@login_required
@require_POST
def purchase_request_reject(request, pk):
    if not role_can_access_url(request.user, 'purchase_request_reject'): raise PermissionDenied
    pr = get_object_or_404(scope_queryset_for_user(PurchaseRequest.objects.all(), request.user), pk=pk)
    form = RejectionForm(request.POST)
    if form.is_valid():
        try: reject_purchase_request(pr, request.user, form.cleaned_data['reason']); messages.success(request, 'Purchase request rejected with reason.')
        except (ValidationError, PermissionDenied) as exc: messages.error(request, str(exc))
    else: messages.error(request, 'A rejection reason of at least five characters is required.')
    return redirect('purchase_request_detail', pk=pk)


@login_required
@require_POST
def purchase_request_processing(request, pk):
    if not role_can_access_url(request.user, 'purchase_request_processing'): raise PermissionDenied
    pr = get_object_or_404(scope_queryset_for_user(PurchaseRequest.objects.all(), request.user), pk=pk)
    try: mark_pr_processing(pr, request.user); messages.success(request, 'PR marked as Purchase Processing.')
    except ValidationError as exc: messages.error(request, str(exc))
    return redirect('purchase_request_detail', pk=pk)


class PurchaseOrderListView(ModuleListView):
    model = PurchaseOrder; title = 'Purchase Orders'; columns = [('po_number', 'PO'), ('purchase_request.request_no', 'PR'), ('vendor', 'Vendor'), ('item.item_no', 'Item'), ('quantity', 'Qty'), ('total_amount', 'Amount'), ('status', 'Status')]; create_url_name = 'purchase_order_create'
    def get_select_related(self): return ['purchase_request', 'item', 'plant']


class PurchaseOrderCreateView(SecureCreateView):
    model = PurchaseOrder; form_class = PurchaseOrderForm; title = 'Create Purchase Order from Approved PR'; success_url_name = 'purchase_order_list'

    def get_initial(self):
        initial = super().get_initial()
        pr_id = self.request.GET.get('pr')
        if pr_id:
            permitted = scope_queryset_for_user(
                PurchaseRequest.objects.filter(status__in=['APPROVED', 'PROCESSING']),
                self.request.user,
            ).filter(pk=pr_id).first()
            if permitted:
                initial.update({
                    'purchase_request': permitted,
                    'quantity': permitted.quantity,
                    'description': permitted.purpose,
                })
        return initial

    def form_valid(self, form):
        try:
            with transaction.atomic():
                pr = PurchaseRequest.objects.select_for_update().select_related('plant', 'item').get(pk=form.instance.purchase_request_id)
                assert_object_access(self.request.user, pr)
                if pr.status not in ['APPROVED', 'PROCESSING']:
                    raise ValidationError('Only an approved PR can be converted to a PO.')
                if pr.purchase_orders.exists():
                    raise ValidationError('A purchase order has already been created for this PR.')
                if form.instance.quantity != pr.quantity:
                    raise ValidationError('PO quantity must equal the approved PR quantity in this one-stage workflow.')
                form.instance.purchase_request = pr
                form.instance.plant = pr.plant
                form.instance.item = pr.item
                form.instance.po_type = form.instance.expected_po_type
                form.instance.status = 'DRAFT'
                form.instance.created_by = self.request.user
                response = super().form_valid(form)
                previous = pr.status
                pr.status = 'PO_CREATED'
                pr.save(update_fields=['status', 'updated_at'])
                PurchaseRequestApprovalHistory.objects.create(
                    purchase_request=pr, action='PO_CREATED', previous_status=previous,
                    new_status='PO_CREATED', performed_by=self.request.user, remarks=self.object.po_number,
                )
            return response
        except (ValidationError, IntegrityError) as exc:
            message = '; '.join(exc.messages) if isinstance(exc, ValidationError) else 'This PR has already been converted to a PO.'
            form.add_error('purchase_request', message)
            return self.form_invalid(form)


@login_required
@require_POST
def purchase_order_action(request, pk, action):
    if not role_can_access_url(request.user, 'purchase_order_action'):
        raise PermissionDenied
    po = get_object_or_404(scope_queryset_for_user(PurchaseOrder.objects.select_related('purchase_request'), request.user), pk=pk)
    if action == 'release':
        if po.status != 'DRAFT':
            messages.error(request, 'Only Draft POs can be released.')
        else:
            po.status = 'RELEASED'; po.save(update_fields=['status', 'updated_at']); messages.success(request, 'Purchase order released.')
    elif action == 'close':
        if po.po_type == 'MATERIAL' and po.status not in {'RECEIVED', 'INVOICED'}:
            messages.error(request, 'Material PO must be fully received before closure.')
        elif po.po_type == 'SERVICE' and po.service_entries.filter(status='ACCEPTED').aggregate(total=Sum('quantity'))['total'] != po.quantity:
            messages.error(request, 'Service PO must be fully accepted before closure.')
        elif po.vendor_invoices.exists() and po.vendor_invoices.exclude(payment_status='RELEASED').exists():
            messages.error(request, 'All vendor invoices must be payment-released before PO closure.')
        else:
            po.status='CLOSED'; po.save(update_fields=['status','updated_at'])
            pr=po.purchase_request; pr.status='CLOSED'; pr.save(update_fields=['status','updated_at']); messages.success(request,'Purchase order closed.')
    elif action == 'cancel':
        if po.goods_receipts.filter(status='POSTED').exists() or po.service_entries.filter(status='ACCEPTED').exists() or po.vendor_invoices.exists():
            messages.error(request, 'PO cannot be cancelled after receipt/service/invoice activity.')
        elif po.status not in {'DRAFT','RELEASED'}:
            messages.error(request, 'Only Draft or Released POs can be cancelled.')
        else:
            po.status='CANCELLED'; po.save(update_fields=['status','updated_at']); messages.success(request,'Purchase order cancelled.')
    else:
        raise Http404
    return redirect('purchase_order_list')


class ServiceEntryListView(ModuleListView):
    model = ServiceEntrySheet; title = 'Service Entry Sheets'
    columns = [('entry_no','SES'),('purchase_order.po_number','PO'),('work_order_operation.sequence','Operation'),('service_date','Date'),('amount','Amount'),('status','Status')]
    create_url_name = 'service_entry_create'
    def get_select_related(self): return ['purchase_order','purchase_order__plant','work_order_operation']


class ServiceEntryCreateView(SecureCreateView):
    model = ServiceEntrySheet; form_class = ServiceEntrySheetForm; title = 'Create Service Entry Sheet'; success_url_name = 'service_entry_list'
    def form_valid(self, form):
        form.instance.created_by = self.request.user
        po = form.instance.purchase_order
        if po.po_type != 'SERVICE':
            form.add_error('purchase_order','Select a Service PO.'); return self.form_invalid(form)
        if not form.instance.work_order_operation_id and po.purchase_request.source_operation_id:
            form.instance.work_order_operation = po.purchase_request.source_operation
        return super().form_valid(form)


@login_required
@require_POST
def service_entry_action(request, pk, action):
    if not role_can_access_url(request.user, 'service_entry_action'): raise PermissionDenied
    entry = get_object_or_404(scope_queryset_for_user(ServiceEntrySheet.objects.select_related('purchase_order'), request.user), pk=pk)
    if action == 'accept':
        if entry.status != 'DRAFT': messages.error(request,'Only draft service entries can be accepted.')
        else:
            entry.status='ACCEPTED'; entry.accepted_by=request.user; entry.accepted_at=timezone.now(); entry.acceptance_notes=request.POST.get('notes','').strip(); entry.save(); messages.success(request,'Service entry accepted and available for work-order actual costing.')
    elif action == 'reject':
        if entry.status != 'DRAFT': messages.error(request,'Only draft service entries can be rejected.')
        else:
            entry.status='REJECTED'; entry.acceptance_notes=request.POST.get('notes','').strip(); entry.save(); messages.success(request,'Service entry rejected.')
    else: raise Http404
    return redirect('service_entry_list')


class GoodsReceiptListView(ModuleListView):
    model = GoodsReceipt; title = 'MIGO / Goods Receipts'; columns = [('receipt_no', 'Receipt'), ('purchase_order.po_number', 'PO'), ('item.item_no', 'Item'), ('quantity_received', 'Qty'), ('received_location', 'Location'), ('status', 'Status')]; create_url_name = 'goods_receipt_create'
    def get_select_related(self): return ['purchase_order', 'item', 'plant', 'received_location']
class GoodsReceiptCreateView(SecureCreateView):
    model = GoodsReceipt; form_class = GoodsReceiptForm; title = 'Create Goods Receipt'; success_url_name = 'goods_receipt_list'
    def form_valid(self, form): form.instance.plant = form.instance.purchase_order.plant; form.instance.item = form.instance.purchase_order.item; return super().form_valid(form)
@login_required
@require_POST
def goods_receipt_post(request, pk):
    if not role_can_access_url(request.user, 'goods_receipt_post'): raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(GoodsReceipt.objects.all(), request.user), pk=pk)
    try: post_goods_receipt(obj, request.user); messages.success(request, 'Goods receipt posted and stock updated atomically.')
    except ValidationError as exc: messages.error(request, str(exc))
    return redirect('goods_receipt_list')


class StockTransferListView(ModuleListView):
    model = StockTransfer; title = 'Stock Transfers'; columns = [('transfer_no', 'Transfer'), ('item.item_no', 'Item'), ('from_location', 'From'), ('to_location', 'To'), ('quantity', 'Qty'), ('status', 'Status')]; create_url_name = 'stock_transfer_create'
    def get_select_related(self): return ['item', 'plant', 'from_location', 'to_location']
class StockTransferCreateView(SecureCreateView):
    model = StockTransfer; form_class = StockTransferForm; title = 'Create Stock Transfer'; success_url_name = 'stock_transfer_list'
    def form_valid(self, form):
        form.instance.plant = form.instance.item.plant
        return super().form_valid(form)
@login_required
@require_POST
def stock_transfer_post(request, pk):
    if not role_can_access_url(request.user, 'stock_transfer_post'): raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(StockTransfer.objects.all(), request.user), pk=pk)
    try: post_stock_transfer(obj, request.user); messages.success(request, 'Stock transfer posted atomically.')
    except ValidationError as exc: messages.error(request, str(exc))
    return redirect('stock_transfer_list')


class MaterialIssueListView(ModuleListView):
    model = MaterialIssue; title = 'Material Issues'; columns = [('issue_no', 'Issue'), ('item.item_no', 'Item'), ('work_order.wo_number', 'Work Order'), ('source_location', 'Location'), ('quantity_issued', 'Qty'), ('status', 'Status')]; create_url_name = 'material_issue_create'
    def get_select_related(self): return ['item', 'plant', 'work_order', 'source_location']
class MaterialIssueCreateView(SecureCreateView):
    model = MaterialIssue; form_class = MaterialIssueForm; title = 'Create Material Issue'; success_url_name = 'material_issue_list'
    def get_initial(self):
        initial = super().get_initial()
        wo_id = self.request.GET.get('work_order')
        if wo_id:
            wo = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user).filter(pk=wo_id).first()
            if wo:
                initial.update({'work_order': wo, 'plant': wo.plant})
        return initial
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        wo_id = self.request.POST.get('work_order') or self.request.GET.get('work_order')
        if wo_id and 'operation' in form.fields:
            wo = scope_queryset_for_user(WorkOrder.objects.all(), self.request.user).filter(pk=wo_id).first()
            if wo:
                form.fields['operation'].queryset = wo.operations.all()
        return form
    def form_valid(self, form):
        if form.instance.work_order_id and form.instance.work_order.technical_lock:
            form.add_error('work_order', 'TECO work orders are technically locked and cannot receive new material issues.')
            return self.form_invalid(form)
        if form.instance.operation_id and form.instance.operation.work_order_id != form.instance.work_order_id:
            form.add_error('operation', 'Operation must belong to the selected work order.')
            return self.form_invalid(form)
        form.instance.plant = form.instance.item.plant
        return super().form_valid(form)
@login_required
@require_POST
def material_issue_post(request, pk):
    if not role_can_access_url(request.user, 'material_issue_post'): raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(MaterialIssue.objects.all(), request.user), pk=pk)
    try: post_material_issue(obj, request.user); messages.success(request, 'Material issue posted and negative stock prevented.')
    except ValidationError as exc: messages.error(request, str(exc))
    return redirect('material_issue_list')


class VendorInvoiceListView(ModuleListView):
    model = VendorInvoice; title = 'Vendor Invoices'; columns = [('invoice_no', 'Invoice'), ('purchase_order.po_number', 'PO'), ('amount', 'Amount'), ('accounts_clearance_status', 'Clearance'), ('payment_status', 'Payment')]
    create_url_name = 'vendor_invoice_create'
    def get_select_related(self): return ['purchase_order', 'goods_receipt']


class VendorInvoiceCreateView(SecureCreateView):
    model = VendorInvoice; form_class = VendorInvoiceForm; title = 'Create Vendor Invoice'; success_url_name = 'vendor_invoice_list'

    def form_valid(self, form):
        form.instance.accounts_clearance_status = 'PENDING'
        form.instance.payment_status = 'NOT_RELEASED'
        return super().form_valid(form)


@login_required
@require_POST
def vendor_invoice_action(request, pk, action):
    if not role_can_access_url(request.user, 'vendor_invoice_action'):
        raise PermissionDenied
    invoice = get_object_or_404(scope_queryset_for_user(VendorInvoice.objects.select_related('purchase_order'), request.user), pk=pk)
    notes = request.POST.get('notes', '').strip()
    try:
        if action == 'hold':
            if invoice.accounts_clearance_status not in {'PENDING', 'CLEARED'}:
                raise ValidationError('Only Pending/Cleared invoices can be put on hold.')
            invoice.accounts_clearance_status = 'HOLD'; invoice.clearance_notes = notes
        elif action == 'clear':
            if invoice.accounts_clearance_status not in {'PENDING', 'HOLD'}:
                raise ValidationError('Only Pending/Hold invoices can be cleared.')
            invoice.accounts_clearance_status = 'CLEARED'; invoice.cleared_by = request.user; invoice.clearance_notes = notes
        elif action == 'release_payment':
            if invoice.accounts_clearance_status != 'CLEARED':
                raise ValidationError('Accounts clearance is required before payment release.')
            invoice.payment_status = 'RELEASED'
            po = invoice.purchase_order
            if not po.vendor_invoices.exclude(pk=invoice.pk).exclude(payment_status='RELEASED').exists():
                po.status = 'INVOICED'; po.save(update_fields=['status','updated_at'])
        else:
            raise Http404
        invoice.save()
        messages.success(request, 'Vendor invoice status updated.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('vendor_invoice_list')


# -----------------------------------------------------------------------------
# Utility, CAPEX and disposal (Observations 19, 21, 22)
# -----------------------------------------------------------------------------

class MeterListView(ModuleListView):
    model = UtilityMeter; title = 'Utility and Electrical Meters'; columns = [('meter_code', 'Meter'), ('name', 'Name'), ('plant.code', 'Plant'), ('utility_type', 'Utility'), ('unit', 'Unit'), ('target_per_day', 'Daily Target'), ('active', 'Active')]; create_url_name = 'meter_create'
    def get_select_related(self): return ['plant']
class MeterCreateView(SecureCreateView): model = UtilityMeter; form_class = UtilityMeterForm; title = 'Create Utility Meter'; success_url_name = 'meter_list'
class MeterReadingCreateView(SecureCreateView):
    model = MeterReading; form_class = MeterReadingForm; title = 'Add Meter Reading'; success_url_name = 'utility_dashboard'
    def form_valid(self, form): form.instance.recorded_by = self.request.user; return super().form_valid(form)


class CAPEXListView(ModuleListView):
    model = CAPEXProposal; title = 'CAPEX Proposals'; columns = [('proposal_no', 'Proposal'), ('title', 'Title'), ('plant.code', 'Plant'), ('budget_year', 'Budget Year'), ('amount', 'Amount'), ('assigned_approver.username', 'Approver'), ('status', 'Status')]; create_url_name = 'capex_create'
    def get_select_related(self): return ['plant', 'work_centre', 'assigned_approver']
class CAPEXCreateView(SecureCreateView):
    model = CAPEXProposal; form_class = CAPEXProposalForm; title = 'Create CAPEX Proposal'; success_url_name = 'capex_list'
    def form_valid(self, form): form.instance.requested_by = self.request.user; return super().form_valid(form)
@login_required
@require_POST
def capex_action(request, pk, action):
    if not role_can_access_url(request.user, 'capex_action'):
        raise PermissionDenied
    try:
        with transaction.atomic():
            queryset = scope_queryset_for_user(CAPEXProposal.objects.select_for_update(), request.user)
            obj = get_object_or_404(queryset, pk=pk)
            previous = obj.status
            remarks = request.POST.get('remarks', '').strip()
            if action == 'submit' and obj.status in ['DRAFT', 'REJECTED'] and obj.requested_by_id == request.user.id:
                obj.status = 'SUBMITTED'
                obj.approval_remarks = ''
            elif action == 'approve' and obj.status == 'SUBMITTED' and (obj.assigned_approver_id == request.user.id or is_admin(request.user)):
                if obj.requested_by_id == request.user.id:
                    raise PermissionDenied('Self-approval is not allowed.')
                if len(remarks) < 3:
                    raise ValidationError('Approval remarks are required.')
                obj.status = 'APPROVED'
                obj.approved_by = request.user
                obj.approved_at = timezone.now()
                obj.approval_remarks = remarks
            elif action == 'reject' and obj.status == 'SUBMITTED' and (obj.assigned_approver_id == request.user.id or is_admin(request.user)):
                if len(remarks) < 5:
                    raise ValidationError('A rejection reason is required.')
                obj.status = 'REJECTED'
                obj.approval_remarks = remarks
            else:
                raise ValidationError('Invalid CAPEX transition or permission.')
            obj.save()
            log_action(request, 'STATUS', obj, previous={'status': previous}, remarks=f'{action}: {remarks}')
        messages.success(request, f'{obj.proposal_no} is now {obj.get_status_display()}.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('capex_list')


@login_required
def capex_pdf(request, pk):
    if not role_can_access_url(request.user, 'capex_pdf'): raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(CAPEXProposal.objects.all(), request.user), pk=pk)
    return safe_pdf_response('CAPEX Approval Note', [('Proposal', obj.proposal_no), ('Status', obj.get_status_display()), ('Plant', obj.plant), ('Work Centre', obj.work_centre), ('Budget Year', obj.budget_year), ('Cost Centre', obj.cost_centre), ('Title', obj.title), ('Justification', obj.justification), ('Business Benefit', obj.business_benefit), ('Amount', obj.amount), ('Requested By', obj.requested_by), ('Approved By', obj.approved_by), ('Remarks', obj.approval_remarks)], f'{obj.proposal_no}.pdf')


class DisposalListView(ModuleListView):
    model = AssetDisposalRequest; title = 'Asset Disposal Approval'; columns = [('disposal_no', 'Disposal'), ('asset.asset_id', 'Asset'), ('plant.code', 'Plant'), ('disposal_method', 'Method'), ('estimated_value', 'Value'), ('assigned_approver.username', 'Approver'), ('status', 'Status')]; create_url_name = 'disposal_create'
    def get_select_related(self): return ['asset', 'plant', 'assigned_approver']
class DisposalCreateView(SecureCreateView):
    model = AssetDisposalRequest; form_class = AssetDisposalForm; title = 'Create Asset Disposal Request'; success_url_name = 'disposal_list'
    def form_valid(self, form): form.instance.requested_by = self.request.user; form.instance.plant = form.instance.asset.plant; return super().form_valid(form)
@login_required
@require_POST
def disposal_action(request, pk, action):
    if not role_can_access_url(request.user, 'disposal_action'):
        raise PermissionDenied
    try:
        with transaction.atomic():
            queryset = scope_queryset_for_user(AssetDisposalRequest.objects.select_for_update().select_related('asset'), request.user)
            obj = get_object_or_404(queryset, pk=pk)
            previous = obj.status
            remarks = request.POST.get('remarks', '').strip()
            if action == 'submit' and obj.status in ['DRAFT', 'REJECTED'] and obj.requested_by_id == request.user.id:
                obj.status = 'SUBMITTED'
                obj.approval_remarks = ''
                obj.asset.status = 'DISPOSAL_PENDING'
                obj.asset.save(update_fields=['status', 'updated_at'])
            elif action == 'approve' and obj.status == 'SUBMITTED' and (obj.assigned_approver_id == request.user.id or is_admin(request.user)):
                if obj.requested_by_id == request.user.id:
                    raise PermissionDenied('Self-approval is not allowed.')
                if len(remarks) < 3:
                    raise ValidationError('Approval remarks are required.')
                obj.status = 'APPROVED'
                obj.approved_by = request.user
                obj.approved_at = timezone.now()
                obj.approval_remarks = remarks
            elif action == 'reject' and obj.status == 'SUBMITTED' and (obj.assigned_approver_id == request.user.id or is_admin(request.user)):
                if len(remarks) < 5:
                    raise ValidationError('A rejection reason is required.')
                obj.status = 'REJECTED'
                obj.approval_remarks = remarks
                obj.asset.status = 'RUNNING'
                obj.asset.save(update_fields=['status', 'updated_at'])
            elif action == 'execute' and obj.status == 'APPROVED' and (obj.assigned_approver_id == request.user.id or is_admin(request.user)):
                if len(remarks) < 3:
                    raise ValidationError('Execution remarks are required.')
                obj.status = 'EXECUTED'
                obj.approval_remarks = f'{obj.approval_remarks}\nExecution: {remarks}'.strip()
                obj.asset.status = 'SCRAPPED'
                obj.asset.save(update_fields=['status', 'updated_at'])
            else:
                raise ValidationError('Invalid disposal transition or permission.')
            obj.save()
            log_action(request, 'STATUS', obj, previous={'status': previous}, remarks=f'{action}: {remarks}')
        messages.success(request, f'{obj.disposal_no} is now {obj.get_status_display()}.')
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
    return redirect('disposal_list')


@login_required
def disposal_pdf(request, pk):
    if not role_can_access_url(request.user, 'disposal_pdf'): raise PermissionDenied
    obj = get_object_or_404(scope_queryset_for_user(AssetDisposalRequest.objects.all(), request.user), pk=pk)
    return safe_pdf_response('Asset Disposal Approval Note', [('Disposal No', obj.disposal_no), ('Status', obj.get_status_display()), ('Plant', obj.plant), ('Asset', obj.asset), ('Condition', obj.condition), ('Reason', obj.reason), ('Estimated Value', obj.estimated_value), ('Method', obj.get_disposal_method_display()), ('Requested By', obj.requested_by), ('Approved By', obj.approved_by), ('Remarks', obj.approval_remarks)], f'{obj.disposal_no}.pdf')


class AuditListView(RoleRequiredMixin, ListView):
    model = AuditLog; template_name = 'asset_mgmt/audit_list.html'; context_object_name = 'items'; paginate_by = 100
    def get_queryset(self): return AuditLog.objects.select_related('user').all()
