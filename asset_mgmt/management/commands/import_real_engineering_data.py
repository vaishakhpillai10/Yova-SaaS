"""Import real engineering Excel data into the Engineering SaaS project.

Copy this file to:
asset_mgmt/management/commands/import_real_engineering_data.py

Usage:
    py manage.py import_real_engineering_data --dry-run
    py manage.py import_real_engineering_data
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - shown to user in command line
    load_workbook = None

from asset_mgmt.models import (
    Area,
    Asset,
    AssetCategory,
    AssetSubcategory,
    AssetDisposalRequest,
    Company,
    CostCentre,
    Department,
    FunctionalLocation,
    Plant,
    UserAssignment,
    WorkCentre,
)
from asset_mgmt.roles import ALL_ROLES, ROLE_HOD

User = get_user_model()


@dataclass
class ImportStats:
    created: dict[str, int] = field(default_factory=dict)
    updated: dict[str, int] = field(default_factory=dict)
    skipped: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def inc(self, bucket: str, name: str, amount: int = 1):
        target = getattr(self, bucket)
        target[name] = target.get(name, 0) + amount

    def warn(self, message: str):
        if len(self.warnings) < 80:
            self.warnings.append(message)


class Command(BaseCommand):
    help = 'Safely import real engineering data from Engg.xlsx and Asset Disposal Form Excel files.'

    def add_arguments(self, parser):
        parser.add_argument('--engg', default=None, help='Path to Engg.xlsx. Defaults to data_imports/Engg.xlsx')
        parser.add_argument('--disposal', default=None, help='Path to Asset Disposal Form (1).xlsx. Defaults to data_imports/Asset Disposal Form (1).xlsx')
        parser.add_argument('--dry-run', action='store_true', help='Validate and preview import without saving anything.')
        parser.add_argument('--skip-disposal', action='store_true', help='Skip Asset Disposal Form import.')
        parser.add_argument('--password', default='RealData@123', help='Password for generated responsible/HOD users.')

    def handle(self, *args, **options):
        if load_workbook is None:
            raise CommandError('openpyxl is not installed. Run: pip install openpyxl')

        self.stats = ImportStats()
        self.generated_password = options['password']
        base_dir = Path(settings.BASE_DIR)
        engg_path = Path(options['engg'] or base_dir / 'data_imports' / 'Engg.xlsx')
        disposal_path = Path(options['disposal'] or base_dir / 'data_imports' / 'Asset Disposal Form (1).xlsx')

        if not engg_path.exists():
            raise CommandError(f'Engg.xlsx not found: {engg_path}')
        if not options['skip_disposal'] and not disposal_path.exists():
            raise CommandError(f'Asset Disposal Form not found: {disposal_path}')

        self.stdout.write(self.style.WARNING('DRY RUN: no database changes will be saved.') if options['dry_run'] else self.style.WARNING('LIVE IMPORT: database changes will be saved.'))

        with transaction.atomic():
            self.ensure_roles()
            self.company = self.get_company()
            engg_wb = load_workbook(engg_path, read_only=True, data_only=True)
            self.import_person_responsible(engg_wb)
            self.import_work_centres(engg_wb)
            self.import_functional_locations(engg_wb)
            self.import_equipment_master(engg_wb)
            engg_wb.close()

            if not options['skip_disposal']:
                disposal_wb = load_workbook(disposal_path, read_only=True, data_only=True)
                self.import_asset_disposal(disposal_wb)
                disposal_wb.close()

            if options['dry_run']:
                transaction.set_rollback(True)

        self.print_summary(options['dry_run'])

    # ---------- common helpers ----------

    def clean(self, value):
        if value is None:
            return ''
        if isinstance(value, str):
            return ' '.join(value.replace('\n', ' ').split())
        return value

    def as_text(self, value):
        value = self.clean(value)
        if value == '':
            return ''
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def safe_code(self, value, fallback='GEN', max_len=30):
        value = self.as_text(value)
        value = re.sub(r'[^A-Za-z0-9_-]+', '_', value).strip('_').upper()
        return (value or fallback)[:max_len]

    def parse_decimal(self, value, default='0'):
        value = self.clean(value)
        if value == '':
            return Decimal(default)
        try:
            return Decimal(str(value).replace(',', '').strip())
        except (InvalidOperation, AttributeError):
            return Decimal(default)

    def parse_date(self, value):
        value = self.clean(value)
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if not value:
            return None
        text = self.as_text(value)
        for fmt in ('%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
        return None

    def ensure_roles(self):
        for role in ALL_ROLES:
            Group.objects.get_or_create(name=role)

    def get_company(self):
        company, created = Company.objects.get_or_create(
            code='IIL',
            defaults={'name': 'Insecticides (India) Limited'},
        )
        self.stats.inc('created' if created else 'updated', 'Company')
        return company

    def get_plant(self, plant_code, plant_name=''):
        code = self.safe_code(plant_code, fallback='PLANT', max_len=30)
        if not code:
            code = 'PLANT'
        name = self.as_text(plant_name) or f'Plant {code}'
        plant, created = Plant.objects.get_or_create(
            company=self.company,
            code=code,
            defaults={'name': name, 'active': True},
        )
        if not created and plant.name.startswith('Plant ') and name:
            plant.name = name[:150]
            plant.save(update_fields=['name', 'updated_at'])
            self.stats.inc('updated', 'Plant')
        else:
            self.stats.inc('created' if created else 'updated', 'Plant')
        return plant

    def discipline_from(self, code='', description=''):
        text = f'{code} {description}'.upper()
        if 'ELECT' in text:
            return 'ELECTRICAL'
        if 'INST' in text:
            return 'INSTRUMENT'
        if 'CIVIL' in text:
            return 'CIVIL'
        if 'UTIL' in text or 'BOILER' in text:
            return 'UTILITY'
        if 'EHS' in text or 'SAFETY' in text:
            return 'SAFETY'
        if 'IT' in text:
            return 'IT'
        if 'STORE' in text:
            return 'STORE'
        if 'MM' in text or 'PURCHASE' in text or 'MATERIAL' in text:
            return 'MM'
        if 'PROD' in text:
            return 'PRODUCTION'
        if 'MECH' in text or 'MEC' in text:
            return 'MECHANICAL'
        return 'OTHER'

    def direction_from_label(self, label):
        token = self.as_text(label).split('-')[-1].upper()
        mapping = {'NX': 'N', 'SX': 'S', 'EX': 'E', 'WX': 'W', 'NE': 'NE', 'NW': 'NW', 'SE': 'SE', 'SW': 'SW', 'N': 'N', 'S': 'S', 'E': 'E', 'W': 'W'}
        return mapping.get(token, 'C')

    def floor_from_label(self, label):
        parts = self.as_text(label).split('-')
        for token in reversed(parts):
            token = token.strip().upper()
            if re.fullmatch(r'\d+', token):
                return token[:20]
        return ''


    def get_department(self, plant, code, name=''):
        code = self.safe_code(code or name, fallback='GEN', max_len=30)
        name = self.as_text(name) or code
        dept, _ = Department.objects.get_or_create(plant=plant, code=code, defaults={'name': name[:150], 'active': True})
        return dept

    def get_cost_centre(self, plant, code, name=''):
        code = self.safe_code(code, fallback='', max_len=40)
        if not code:
            return None
        centre, _ = CostCentre.objects.get_or_create(plant=plant, code=code, defaults={'name': (self.as_text(name) or code)[:150], 'active': True})
        return centre

    # ---------- import stages ----------

    def import_person_responsible(self, wb):
        self.responsible_by_plant_code = {}
        if 'Person Responsible' not in wb.sheetnames:
            self.stats.warn('Person Responsible sheet missing.')
            return
        ws = wb['Person Responsible']
        for idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            plant_code = self.as_text(row[0] if len(row) > 0 else '')
            plant_desc = self.as_text(row[1] if len(row) > 1 else '')
            responsible_code = self.safe_code(row[2] if len(row) > 2 else '', max_len=30)
            department = self.as_text(row[3] if len(row) > 3 else '')
            if not plant_code or not responsible_code:
                self.stats.inc('skipped', 'Person Responsible')
                continue
            self.get_plant(plant_code, plant_desc)
            self.responsible_by_plant_code[(plant_code, responsible_code)] = department

    def import_work_centres(self, wb):
        if 'Work Center' not in wb.sheetnames:
            self.stats.warn('Work Center sheet missing.')
            return
        ws = wb['Work Center']
        group = Group.objects.get(name=ROLE_HOD)
        for idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            plant_name = self.as_text(row[0] if len(row) > 0 else '')
            plant_code = self.as_text(row[1] if len(row) > 1 else '')
            wc_code = self.safe_code(row[2] if len(row) > 2 else '', max_len=30)
            wc_name = self.as_text(row[3] if len(row) > 3 else '')
            responsible_code = self.safe_code(row[4] if len(row) > 4 else '', max_len=30)
            if not plant_code or not wc_code:
                self.stats.inc('skipped', 'WorkCentre')
                continue
            plant = self.get_plant(plant_code, plant_name)
            self.get_department(plant, responsible_code or wc_code, wc_name or responsible_code)
            if len(row) > 13:
                self.get_cost_centre(plant, row[13], wc_name or wc_code)
            wc, created = WorkCentre.objects.update_or_create(
                plant=plant,
                code=wc_code,
                defaults={
                    'name': (wc_name or wc_code)[:120],
                    'discipline': self.discipline_from(wc_code, wc_name),
                    'active': True,
                },
            )
            self.stats.inc('created' if created else 'updated', 'WorkCentre')
            if responsible_code:
                department = self.responsible_by_plant_code.get((plant_code, responsible_code), wc_name or responsible_code)
                username = f'hod_{plant.code.lower()}_{slugify(responsible_code).replace("-", "_")}'[:150]
                user, user_created = User.objects.get_or_create(username=username)
                user.first_name = responsible_code[:30]
                user.last_name = department[:150]
                user.email = f'{username}@example.local'
                user.is_active = True
                user.set_password(self.generated_password)
                user.save()
                user.groups.add(group)
                UserAssignment.objects.update_or_create(
                    user=user,
                    plant=plant,
                    work_centre=wc,
                    defaults={'designation': 'HOD', 'active': True, 'is_primary': False},
                )
                self.stats.inc('created' if user_created else 'updated', 'Generated HOD User')

    def import_functional_locations(self, wb):
        if 'Functional Location' not in wb.sheetnames:
            self.stats.warn('Functional Location sheet missing.')
            return
        ws = wb['Functional Location']
        parent_map = []
        for idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            fl_code = self.as_text(row[1] if len(row) > 1 else '')
            desc = self.as_text(row[4] if len(row) > 4 else '')
            maint_plant = self.as_text(row[6] if len(row) > 6 else '')
            location = self.as_text(row[7] if len(row) > 7 else '')
            plant_section = self.as_text(row[8] if len(row) > 8 else '')
            work_centre_code = self.safe_code(row[14] if len(row) > 14 else '', max_len=30)
            if not fl_code or not maint_plant:
                self.stats.inc('skipped', 'FunctionalLocation')
                continue
            plant = self.get_plant(maint_plant)
            area_code = self.safe_code(plant_section or location or self.derive_area_from_fl(fl_code), max_len=30)
            area_name = location or plant_section or area_code
            area, _ = Area.objects.get_or_create(plant=plant, code=area_code, defaults={'name': area_name[:150]})
            if len(row) > 11:
                self.get_cost_centre(plant, row[11], desc or fl_code)
            defaults = {
                'plant': plant,
                'area': area,
                'name': (desc or fl_code)[:180],
                'floor': self.floor_from_label(fl_code),
                'cardinal_direction': self.direction_from_label(fl_code),
                'description': desc,
            }
            fl, created = FunctionalLocation.objects.update_or_create(code=fl_code[:120], defaults=defaults)
            self.stats.inc('created' if created else 'updated', 'FunctionalLocation')
            parent_code = self.parent_fl_code(fl_code)
            if parent_code:
                parent_map.append((fl.pk, parent_code[:120]))
            if work_centre_code:
                WorkCentre.objects.get_or_create(
                    plant=plant,
                    code=work_centre_code,
                    defaults={'name': work_centre_code, 'discipline': self.discipline_from(work_centre_code), 'active': True},
                )
        for fl_pk, parent_code in parent_map:
            parent = FunctionalLocation.objects.filter(code=parent_code).first()
            if parent:
                FunctionalLocation.objects.filter(pk=fl_pk).exclude(parent=parent).update(parent=parent)

    def derive_area_from_fl(self, fl_code):
        parts = self.as_text(fl_code).split('-')
        return parts[2] if len(parts) >= 3 else 'GEN'

    def parent_fl_code(self, fl_code):
        text = self.as_text(fl_code)
        parts = text.split('-')
        if len(parts) <= 1:
            return ''
        return '-'.join(parts[:-1])

    def import_equipment_master(self, wb):
        if 'Equipment Master' not in wb.sheetnames:
            self.stats.warn('Equipment Master sheet missing.')
            return
        ws = wb['Equipment Master']
        parent_links = []
        category_names = self.equipment_category_names(wb)
        for idx, row in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
            equipment_no = self.as_text(row[1] if len(row) > 1 else '')
            description = self.as_text(row[4] if len(row) > 4 else '')
            if not equipment_no or not description:
                self.stats.inc('skipped', 'Equipment')
                continue
            equipment_category = self.safe_code(row[3] if len(row) > 3 else '', fallback='OTHER', max_len=10)
            maintenance_plant = self.as_text(row[23] if len(row) > 23 else '') or self.as_text(row[32] if len(row) > 32 else '')
            fl_code = self.as_text(row[38] if len(row) > 38 else '')
            if not maintenance_plant and fl_code:
                parts = fl_code.split('-')
                maintenance_plant = parts[1] if len(parts) > 1 else ''
            if not maintenance_plant:
                self.stats.warn(f'Equipment row {idx}: missing plant, skipped equipment {equipment_no}')
                self.stats.inc('skipped', 'Equipment')
                continue
            plant = self.get_plant(maintenance_plant)
            fl = self.get_or_create_functional_location_for_equipment(plant, fl_code, row)
            category_name = category_names.get(equipment_category, f'SAP Category {equipment_category}')
            category, _ = AssetCategory.objects.get_or_create(name=category_name[:120])
            object_type_text = self.as_text(row[6] if len(row) > 6 else '')
            subcategory = None
            if object_type_text:
                sub_code = self.safe_code(object_type_text, fallback='OBJECT', max_len=30)
                subcategory, _ = AssetSubcategory.objects.get_or_create(category=category, code=sub_code, defaults={'name': object_type_text[:120], 'active': True})
            asset_type = self.asset_type_from_category(equipment_category, row[6] if len(row) > 6 else '')
            tag_number = equipment_no[:80]
            department = self.get_department(plant, row[33] if len(row) > 33 else row[34] if len(row) > 34 else '', row[34] if len(row) > 34 else '')
            cost_centre = self.get_cost_centre(plant, row[31] if len(row) > 31 else '', description)
            asset, created = Asset.objects.get_or_create(
                plant=plant,
                tag_number=tag_number,
                defaults={
                    'asset_id': self.unique_asset_id(equipment_no, plant),
                    'asset_type': asset_type,
                    'name': description[:180],
                    'asset_description': description,
                    'category': category,
                    'subcategory': subcategory,
                    'department': department,
                    'responsible_department': department,
                    'cost_centre': cost_centre,
                    'functional_location': fl,
                    'manufacturer': self.as_text(row[17] if len(row) > 17 else '')[:150],
                    'model_number': self.as_text(row[11] if len(row) > 11 else '')[:120],
                    'serial_number': self.as_text(row[22] if len(row) > 22 else '')[:120],
                    'manufacturer_part_number': self.as_text(row[21] if len(row) > 21 else '')[:120],
                    'capacity_rating': self.as_text(row[9] if len(row) > 9 else '')[:150] or self.as_text(row[13] if len(row) > 13 else '')[:150],
                    'material_of_construction': self.as_text(row[12] if len(row) > 12 else '')[:150],
                    'weight': self.parse_decimal(row[7] if len(row) > 7 else '', default='0') or None,
                    'speed': self.parse_decimal(row[10] if len(row) > 10 else '', default='0') or None,
                    'dimensions': self.as_text(row[13] if len(row) > 13 else '')[:150],
                    'purchase_cost': self.parse_decimal(row[15] if len(row) > 15 else '', default='0') or None,
                    'purchase_date': self.parse_date(row[16] if len(row) > 16 else None),
                    'asset_accounting_number': self.as_text(row[29] if len(row) > 29 else '')[:100] or None,
                    'commissioning_date': self.parse_date(row[14] if len(row) > 14 else None) or self.parse_date(row[16] if len(row) > 16 else None),
                    'warranty_start_date': self.parse_date(row[39] if len(row) > 39 else None),
                    'warranty_end_date': self.parse_date(row[40] if len(row) > 40 else None),
                    'criticality': self.criticality_from_abc(row[27] if len(row) > 27 else ''),
                    'criticality_score': self.criticality_score_from_abc(row[27] if len(row) > 27 else ''),
                    'safety_critical': self.criticality_from_abc(row[27] if len(row) > 27 else '') == 'HIGH',
                    'remarks': self.as_text(row[30] if len(row) > 30 else '')[:1000],
                },
            )
            if not created:
                asset.name = description[:180]
                asset.asset_description = description
                asset.asset_type = asset_type
                asset.category = category
                asset.subcategory = subcategory
                asset.department = department
                asset.responsible_department = department
                asset.cost_centre = cost_centre
                asset.functional_location = fl
                asset.manufacturer = self.as_text(row[17] if len(row) > 17 else '')[:150]
                asset.model_number = self.as_text(row[11] if len(row) > 11 else '')[:120]
                asset.serial_number = self.as_text(row[22] if len(row) > 22 else '')[:120]
                asset.manufacturer_part_number = self.as_text(row[21] if len(row) > 21 else '')[:120]
                asset.material_of_construction = self.as_text(row[12] if len(row) > 12 else '')[:150]
                asset.capacity_rating = self.as_text(row[9] if len(row) > 9 else '')[:150] or self.as_text(row[13] if len(row) > 13 else '')[:150]
                asset.purchase_cost = self.parse_decimal(row[15] if len(row) > 15 else '', default='0') or None
                asset.purchase_date = self.parse_date(row[16] if len(row) > 16 else None)
                asset.asset_accounting_number = self.as_text(row[29] if len(row) > 29 else '')[:100] or None
                asset.commissioning_date = self.parse_date(row[14] if len(row) > 14 else None) or self.parse_date(row[16] if len(row) > 16 else None)
                asset.warranty_start_date = self.parse_date(row[39] if len(row) > 39 else None)
                asset.warranty_end_date = self.parse_date(row[40] if len(row) > 40 else None)
                asset.criticality = self.criticality_from_abc(row[27] if len(row) > 27 else '')
                asset.criticality_score = self.criticality_score_from_abc(row[27] if len(row) > 27 else '')
                asset.safety_critical = asset.criticality == 'HIGH'
                asset.remarks = self.as_text(row[30] if len(row) > 30 else '')[:1000]
                asset.save()
            self.stats.inc('created' if created else 'updated', 'Equipment/Asset')
            superior_equipment = self.as_text(row[37] if len(row) > 37 else '')
            if superior_equipment:
                parent_links.append((asset.pk, plant.pk, superior_equipment[:80]))

        for asset_pk, plant_pk, superior_tag in parent_links:
            parent = Asset.objects.filter(plant_id=plant_pk, tag_number=superior_tag).first() or Asset.objects.filter(asset_id=superior_tag).first()
            if parent and parent.pk != asset_pk:
                Asset.objects.filter(pk=asset_pk).update(parent_asset=parent)

    def equipment_category_names(self, wb):
        result = {}
        if 'EQ_KDS' not in wb.sheetnames:
            return result
        ws = wb['EQ_KDS']
        for row in ws.iter_rows(min_row=4, values_only=True):
            code = self.safe_code(row[0] if len(row) > 0 else '', max_len=10)
            desc = self.as_text(row[1] if len(row) > 1 else '')
            if code and desc:
                result[code] = desc
        return result

    def asset_type_from_category(self, category_code, object_type=''):
        code = self.safe_code(category_code, max_len=10)
        object_text = self.as_text(object_type).upper()
        if code == 'E':
            return 'ELECTRICAL'
        if code == 'I':
            return 'INSTRUMENT'
        if code == 'R':
            if 'PUMP' in object_text:
                return 'PUMP'
            return 'PUMP'
        if code == 'V' or code == 'G':
            return 'UTILITY'
        if code == 'S':
            return 'OTHER'
        if code == 'C':
            return 'OTHER'
        return 'OTHER'

    def criticality_from_abc(self, value):
        value = self.as_text(value).upper()
        if value == 'A':
            return 'HIGH'
        if value == 'C':
            return 'LOW'
        return 'MEDIUM'

    def criticality_score_from_abc(self, value):
        value = self.as_text(value).upper()
        if value == 'A':
            return 90
        if value == 'B':
            return 60
        if value == 'C':
            return 30
        return 50

    def unique_asset_id(self, equipment_no, plant):
        base = self.as_text(equipment_no)[:80] or f'ASSET-{plant.code}'
        if not Asset.objects.filter(asset_id=base).exists():
            return base
        alt = f'{base[:60]}-{plant.code}'[:80]
        if not Asset.objects.filter(asset_id=alt).exists():
            return alt
        n = Asset.objects.filter(asset_id__startswith=base[:50]).count() + 1
        return f'{base[:50]}-{plant.code}-{n}'[:80]

    def get_or_create_functional_location_for_equipment(self, plant, fl_code, row):
        fl_code = self.as_text(fl_code)[:120]
        if fl_code:
            found = FunctionalLocation.objects.filter(code=fl_code).first()
            if found:
                return found
        area_code = self.safe_code(row[25] if len(row) > 25 else '', fallback='GEN', max_len=30)
        area, _ = Area.objects.get_or_create(plant=plant, code=area_code, defaults={'name': area_code})
        code = fl_code or f'{plant.company.code}-{plant.code}-{area.code}-GF-C-GENERAL'
        fl, _ = FunctionalLocation.objects.get_or_create(
            code=code[:120],
            defaults={'plant': plant, 'area': area, 'name': code[:180], 'floor': self.floor_from_label(code), 'cardinal_direction': self.direction_from_label(code)},
        )
        return fl

    def import_asset_disposal(self, wb):
        ws = wb[wb.sheetnames[0]]
        asset_name = self.as_text(ws['B3'].value) or 'Asset for Disposal'
        asset_category = self.as_text(ws['B4'].value) or 'Disposal Asset'
        manufacturer = self.as_text(ws['B5'].value)
        moc = self.as_text(ws['B6'].value)
        capacity = self.as_text(ws['B7'].value)
        year_purchase = self.as_text(ws['B8'].value)
        asset_code = self.as_text(ws['B9'].value)
        objective = self.as_text(ws['B2'].value)
        proposed_1 = self.as_text(ws['B19'].value)
        proposed_2 = self.as_text(ws['B20'].value)
        proposed_3 = self.as_text(ws['B21'].value)
        justification = self.as_text(ws['B22'].value)
        benefit = self.as_text(ws['B24'].value)

        plant = Plant.objects.filter(code='1005').first() or Plant.objects.order_by('code').first() or self.get_plant('1005', 'Technical Plant')
        area, _ = Area.objects.get_or_create(plant=plant, code='BG', defaults={'name': 'B & G Plant'})
        fl, _ = FunctionalLocation.objects.get_or_create(
            code=f'{plant.company.code}-{plant.code}-BG-GF-C-DISPOSAL'[:120],
            defaults={'plant': plant, 'area': area, 'name': 'B & G Disposal Area', 'floor': 'GF', 'cardinal_direction': 'C'},
        )
        category, _ = AssetCategory.objects.get_or_create(name=asset_category[:120])
        asset = Asset.objects.filter(asset_id=asset_code).first() or Asset.objects.filter(plant=plant, tag_number=asset_code[:80]).first()
        if asset:
            created_asset = False
            asset.name = asset_name[:180]
            asset.asset_description = asset_name
            asset.category = category
            asset.functional_location = fl
            asset.manufacturer = manufacturer[:150]
            asset.material_of_construction = moc[:150]
            asset.capacity_rating = capacity[:150]
            asset.remarks = f'Year of purchase: {year_purchase}'.strip()[:1000]
            asset.save()
        else:
            asset = Asset.objects.create(
                plant=plant,
                asset_id=asset_code[:80] or self.unique_asset_id(asset_name, plant),
                tag_number=(asset_code or slugify(asset_name) or 'DISPOSAL-ASSET')[:80],
                asset_type='OTHER',
                name=asset_name[:180],
                asset_description=asset_name,
                category=category,
                functional_location=fl,
                manufacturer=manufacturer[:150],
                material_of_construction=moc[:150],
                capacity_rating=capacity[:150],
                criticality='MEDIUM',
                remarks=f'Year of purchase: {year_purchase}'.strip()[:1000],
            )
            created_asset = True
        self.stats.inc('created' if created_asset else 'updated', 'Disposal Asset')

        requested_by = User.objects.filter(username='maintenance_engineer').first() or User.objects.filter(is_superuser=True).first()
        approver = User.objects.filter(username='maintenance_hod').first() or User.objects.filter(is_superuser=True).first()
        if not requested_by or not approver:
            raise CommandError('Could not find maintenance_engineer/maintenance_hod or a superuser for disposal workflow.')

        reason = '\n'.join(part for part in [
            f'Objective: {objective}',
            f'Justification: {justification}',
            f'Benefit: {benefit}',
            f'Proposed option 1: {proposed_1}',
            f'Proposed option 2: {proposed_2}',
            f'Proposed option 3: {proposed_3}',
        ] if part.split(': ', 1)[-1])
        disposal, created = AssetDisposalRequest.objects.get_or_create(
            asset=asset,
            status='SUBMITTED',
            defaults={
                'plant': plant,
                'condition': 'Not utilised from 2020-21 as per proposal note.',
                'reason': reason or 'Asset disposal imported from proposal note.',
                'estimated_value': Decimal('0'),
                'disposal_method': 'SCRAP',
                'requested_by': requested_by,
                'assigned_approver': approver,
            },
        )
        if not created:
            disposal.condition = 'Not utilised from 2020-21 as per proposal note.'
            disposal.reason = reason or disposal.reason
            disposal.disposal_method = 'SCRAP'
            disposal.requested_by = requested_by
            disposal.assigned_approver = approver
            disposal.save()
        self.stats.inc('created' if created else 'updated', 'AssetDisposalRequest')

    def print_summary(self, dry_run):
        self.stdout.write('\n========== REAL DATA IMPORT SUMMARY ==========' )
        for title, bucket in [('Created', self.stats.created), ('Updated/Seen', self.stats.updated), ('Skipped', self.stats.skipped)]:
            self.stdout.write(title + ':')
            if not bucket:
                self.stdout.write('  - None')
            for name, count in sorted(bucket.items()):
                self.stdout.write(f'  - {name}: {count}')
        if self.stats.warnings:
            self.stdout.write('\nWarnings:')
            for message in self.stats.warnings:
                self.stdout.write(f'  - {message}')
        if dry_run:
            self.stdout.write(self.style.WARNING('\nDRY RUN COMPLETE. No data was saved.'))
            self.stdout.write('If the numbers look okay, run: py manage.py import_real_engineering_data')
        else:
            self.stdout.write(self.style.SUCCESS('\nLIVE IMPORT COMPLETE. Data saved successfully.'))
            self.stdout.write(f'Generated HOD users use password: {self.generated_password}')
